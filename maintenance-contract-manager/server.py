import os
import re
import io
import base64
import hashlib
import hmac
import smtplib
import shutil
import tempfile
import threading
import uuid
import zipfile
from datetime import date, datetime, timedelta
from email.message import EmailMessage
from pathlib import Path
from typing import Any
from time import time
from urllib.parse import quote

from organization_links import normalize_organization_pair, resolve_allowed_organization_ids

import fitz
import psycopg
import pytesseract
from fastapi import BackgroundTasks, FastAPI, File, Form, HTTPException, Request, Response, UploadFile
from fastapi.responses import FileResponse, JSONResponse
from fastapi.staticfiles import StaticFiles
from PIL import Image
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from psycopg.rows import dict_row


BASE_DIR = Path(__file__).resolve().parent
ENV_PATH = BASE_DIR / "config.env"
DEFAULT_UPLOAD_DIR = r"\\125.136.150.4\ict\유지보수계약"

DEFAULT_ORGANIZATIONS = [
    ("대표이사 (CEO)", None),
    ("총괄사장 (COO)", "대표이사 (CEO)"),
    ("TDK", "총괄사장 (COO)"),
    ("연구소", "대표이사 (CEO)"),
    ("경영지원실", "총괄사장 (COO)"),
    ("재경 GR", "경영지원실"),
    ("사업지원 GR", "경영지원실"),
    ("인사지원팀", "총괄사장 (COO)"),
    ("안전 GR", "인사지원팀"),
    ("개발구매팀", "총괄사장 (COO)"),
    ("영업개발 GR", "개발구매팀"),
    ("구매 GR", "개발구매팀"),
    ("CS GR", "개발구매팀"),
    ("품질팀", "총괄사장 (COO)"),
    ("평동공장", "총괄사장 (COO)"),
    ("평동 생산총괄팀", "평동공장"),
    ("평동 생산지원 GR", "평동 생산총괄팀"),
    ("평동 자재지원 GR", "평동 생산총괄팀"),
    ("평동 생산1 GR", "평동 생산총괄팀"),
    ("평동 생산2 GR", "평동 생산총괄팀"),
    ("평동 기술 GR", "평동 생산총괄팀"),
    ("하남공장", "총괄사장 (COO)"),
    ("하남 생산총괄팀", "하남공장"),
    ("하남 제품기술 GR", "하남 생산총괄팀"),
    ("하남 생산지원 GR", "하남 생산총괄팀"),
    ("하남 구매 GR", "하남 생산총괄팀"),
    ("하남 자재 GR", "하남 생산총괄팀"),
    ("하남 생산 GR", "하남 생산총괄팀"),
    ("하남 공정기술 GR", "하남 생산총괄팀"),
    ("하남 CS GR", "하남 생산총괄팀"),
    ("하남 품질팀", "하남공장"),
]


def load_env_file() -> dict[str, str]:
    values: dict[str, str] = {}
    if ENV_PATH.exists():
        for raw_line in ENV_PATH.read_text(encoding="utf-8").splitlines():
            line = raw_line.strip()
            if not line or line.startswith("#") or "=" not in line:
                continue
            key, value = line.split("=", 1)
            values[key.strip()] = value.strip().strip('"')
    return values


CONFIG = {
    "DB_HOST": "localhost",
    "DB_PORT": "5432",
    "DB_NAME": "maintenance_db",
    "DB_USER": "maintenance_user",
    "DB_PASSWORD": "Maintenance@2026!",
    "UPLOAD_DIR": DEFAULT_UPLOAD_DIR,
    "TESSERACT_CMD": r"C:\Program Files\Tesseract-OCR\tesseract.exe",
    "SESSION_SECRET": "change-this-session-secret",
    "SMTP_SERVER": "",
    "SMTP_PORT": "25",
    "SENDER_EMAIL": "",
    "SENDER_PASSWORD": "",
    "RECEIVER_EMAIL": "",
    "MAIL_NOTIFY_HOUR": "9",
    "MAIL_NOTIFY_MINUTE": "0",
    **load_env_file(),
}

MAIL_ALERT_DAYS = [60, 30, 14, 7]
MAIL_SCHEDULER_STARTED = False
FULL_ACCESS_ROLES = {"admin", "manager"}
ROLE_PERMISSIONS = {
    "admin": {
        "menu.dashboard.view",
        "menu.schedule.view",
        "menu.stats.view",
        "menu.user.view",
        "menu.user.manage",
        "menu.organization.view",
        "menu.organization.manage",
        "menu.company.view",
        "menu.company.create",
        "menu.company.update",
        "menu.company.delete",
        "menu.contract.view",
        "menu.contract.create",
        "menu.contract.update",
        "menu.contract.delete",
        "menu.pending.view",
        "menu.pending.create",
        "menu.pending.update",
        "menu.pending.delete",
        "menu.asset.view",
        "menu.asset.create",
        "menu.asset.update",
        "menu.asset.delete",
        "menu.worklog.view",
        "menu.worklog.create",
        "menu.worklog.update",
        "menu.worklog.delete",
        "menu.document.view",
        "menu.document.create",
        "menu.document.update",
        "menu.document.delete",
    },
    "manager": {
        "menu.dashboard.view",
        "menu.schedule.view",
        "menu.stats.view",
        "menu.company.view",
        "menu.company.create",
        "menu.company.update",
        "menu.company.delete",
        "menu.contract.view",
        "menu.contract.create",
        "menu.contract.update",
        "menu.contract.delete",
        "menu.pending.view",
        "menu.pending.create",
        "menu.pending.update",
        "menu.pending.delete",
        "menu.asset.view",
        "menu.asset.create",
        "menu.asset.update",
        "menu.asset.delete",
        "menu.worklog.view",
        "menu.worklog.create",
        "menu.worklog.update",
        "menu.worklog.delete",
        "menu.document.view",
        "menu.document.create",
        "menu.document.update",
        "menu.document.delete",
    },
    "team_lead": {
        "menu.dashboard.view",
        "menu.schedule.view",
        "menu.stats.view",
        "menu.company.view",
        "menu.company.create",
        "menu.company.update",
        "menu.contract.view",
        "menu.contract.create",
        "menu.contract.update",
        "menu.pending.view",
        "menu.asset.view",
        "menu.worklog.view",
        "menu.worklog.create",
        "menu.worklog.update",
        "menu.worklog.delete",
        "menu.document.view",
    },
    "team_member": {
        "menu.dashboard.view",
        "menu.schedule.view",
        "menu.stats.view",
        "menu.company.view",
        "menu.company.create",
        "menu.contract.view",
        "menu.contract.create",
        "menu.pending.view",
        "menu.asset.view",
        "menu.worklog.view",
        "menu.worklog.create",
        "menu.document.view",
    },
}

pytesseract.pytesseract.tesseract_cmd = CONFIG["TESSERACT_CMD"]


def conninfo() -> str:
    return (
        f"host={CONFIG['DB_HOST']} "
        f"port={CONFIG['DB_PORT']} "
        f"dbname={CONFIG['DB_NAME']} "
        f"user={CONFIG['DB_USER']} "
        f"password={CONFIG['DB_PASSWORD']}"
    )


def db() -> psycopg.Connection:
    return psycopg.connect(conninfo(), row_factory=dict_row)


def now_iso() -> str:
    return datetime.now().isoformat(timespec="seconds")


def clean_filename(name: str) -> str:
    stem = Path(name).stem or "file"
    suffix = Path(name).suffix
    stem = re.sub(r'[\\/:*?"<>|]+', "_", stem).strip()
    return f"{stem[:80]}{suffix[:20]}"


def normalize_organization_name(value: str) -> str:
    return re.sub(r"[\s\-_·.]+", "", str(value or "").strip().lower())


def department_organization_candidates(parent_name: str, department_name: str) -> list[str]:
    candidates: list[str] = []
    department = str(department_name or "").strip()
    parent = str(parent_name or "").strip()
    if department:
        candidates.append(department)
        match = re.match(r"^(.+)\((하남|평동)\)$", department)
        if match:
            candidates.append(f"{match.group(2).strip()} {match.group(1).strip()}")
    if parent:
        candidates.append(parent)
    result: list[str] = []
    seen: set[str] = set()
    for candidate in candidates:
        key = normalize_organization_name(candidate)
        if key and key not in seen:
            result.append(candidate)
            seen.add(key)
    return result


def row_to_client(row: dict[str, Any]) -> dict[str, Any]:
    converted: dict[str, Any] = {}
    aliases = {
        "business_no": "businessNo",
        "company_id": "companyId",
        "company_name": "companyName",
        "contract_id": "contractId",
        "contract_name": "contractName",
        "manager_user_id": "managerUserId",
        "manager_user_name": "managerUserName",
        "manager_user_email": "managerUserEmail",
        "organization_id": "organizationId",
        "organization_name": "organizationName",
        "user_id": "userId",
        "employee_id": "employeeId",
        "employee_no": "employeeNo",
        "employment_status": "employmentStatus",
        "parent_id": "parentId",
        "sort_order": "sortOrder",
        "is_active": "isActive",
        "start_date": "startDate",
        "end_date": "endDate",
        "billing_cycle": "billingCycle",
        "renewal_date": "renewalDate",
        "alert_days": "alertDays",
        "auto_alert": "autoAlert",
        "work_date": "date",
        "due_date": "dueDate",
        "completed_at": "completedAt",
        "file_name": "fileName",
        "file_path": "filePath",
        "is_primary": "isPrimary",
        "created_at": "createdAt",
        "updated_at": "updatedAt",
    }
    for key, value in row.items():
        client_key = aliases.get(key, key)
        if key == "id" or key.endswith("_id"):
            converted[client_key] = "" if value is None else str(value)
        elif isinstance(value, (datetime, date)):
            converted[client_key] = value.isoformat()
        else:
            converted[client_key] = value
    return converted


def fetch_all(table: str, order_by: str = "id") -> list[dict[str, Any]]:
    with db() as conn:
        with conn.cursor() as cur:
            cur.execute(f"SELECT * FROM {table} ORDER BY {order_by}")
            return [row_to_client(row) for row in cur.fetchall()]


def fetch_worklogs(user: dict[str, Any]) -> list[dict[str, Any]]:
    if user.get("role") in FULL_ACCESS_ROLES:
        return fetch_all("work_logs", "work_date DESC, id DESC")
    organization_ids = [parse_int(value) for value in get_allowed_contract_organization_ids(user)]
    organization_ids = [value for value in organization_ids if value]
    if not organization_ids:
        return []
    with db() as conn:
        with conn.cursor() as cur:
            cur.execute(
                """
                SELECT w.*
                FROM work_logs w
                WHERE EXISTS (
                    SELECT 1
                    FROM app.contract_organizations co
                    WHERE co.contract_id = w.contract_id
                      AND co.organization_id = ANY(%s)
                )
                ORDER BY w.work_date DESC, w.id DESC
                """,
                (organization_ids,),
            )
            return [row_to_client(row) for row in cur.fetchall()]


def scoped_contract_where(user: dict[str, Any], alias: str = "id") -> tuple[str, tuple[Any, ...]]:
    if user.get("role") in FULL_ACCESS_ROLES:
        return "", ()
    organization_ids = [parse_int(value) for value in get_allowed_contract_organization_ids(user)]
    organization_ids = [value for value in organization_ids if value]
    if not organization_ids:
        return " WHERE 1 = 0", ()
    return (
        f" WHERE {alias} IN (SELECT contract_id FROM app.contract_organizations WHERE organization_id = ANY(%s))",
        (organization_ids,),
    )

def fetch_companies(user: dict[str, Any] | None = None) -> list[dict[str, Any]]:
    with db() as conn:
        with conn.cursor() as cur:
            cur.execute(
                """
                SELECT c.*, p.representative
                FROM companies c
                LEFT JOIN app.company_profiles p ON p.company_id = c.id
                ORDER BY c.id
                """,
            )
            companies = [row_to_client(row) for row in cur.fetchall()]
            attach_company_contacts(cur, companies)
            return companies


def fetch_contracts(user: dict[str, Any]) -> list[dict[str, Any]]:
    where_sql, params = scoped_contract_where(user, "c.id")
    with db() as conn:
        with conn.cursor() as cur:
            cur.execute(
                f"""
                SELECT c.*, co.organization_id, o.name AS organization_name, cd.department, cd.manager,
                       cd.manager_user_id,
                       COALESCE(NULLIF(mu.display_name, ''), mu.username) AS manager_user_name,
                       mu.email AS manager_user_email
                FROM contracts c
                LEFT JOIN app.contract_organizations co ON co.contract_id = c.id
                LEFT JOIN app.organizations o ON o.id = co.organization_id
                LEFT JOIN app.contract_details cd ON cd.contract_id = c.id
                LEFT JOIN app.app_users mu ON mu.id = cd.manager_user_id
                {where_sql}
                ORDER BY c.id
                """,
                params,
            )
            return [row_to_client(row) for row in cur.fetchall()]


def fetch_documents(user: dict[str, Any]) -> list[dict[str, Any]]:
    where_sql = ""
    params: tuple[Any, ...] = ()
    if user.get("role") not in FULL_ACCESS_ROLES:
        organization_ids = [parse_int(value) for value in get_allowed_contract_organization_ids(user)]
        organization_ids = [value for value in organization_ids if value]
        if organization_ids:
            where_sql = """
                WHERE (
                    d.contract_id IN (
                        SELECT contract_id FROM app.contract_organizations WHERE organization_id = ANY(%s)
                    )
                    OR d.id IN (
                        SELECT document_id FROM app.document_organizations WHERE organization_id = ANY(%s)
                    )
                )
            """
            params = (organization_ids, organization_ids)
        else:
            where_sql = "WHERE 1 = 0"
    with db() as conn:
        with conn.cursor() as cur:
            cur.execute(
                f"""
                SELECT d.*, doo.organization_id, o.name AS organization_name
                FROM documents d
                LEFT JOIN app.document_organizations doo ON doo.document_id = d.id
                LEFT JOIN app.organizations o ON o.id = doo.organization_id
                {where_sql}
                ORDER BY d.created_at DESC, d.id DESC
                """,
                params,
            )
            return [row_to_client(row) for row in cur.fetchall()]

def fetch_pending_items(user: dict[str, Any] | None = None) -> list[dict[str, Any]]:
    where_sql = ""
    params: tuple[Any, ...] = ()
    if user and user.get("role") not in FULL_ACCESS_ROLES:
        organization_ids = [parse_int(value) for value in get_allowed_contract_organization_ids(user)]
        organization_ids = [value for value in organization_ids if value]
        if organization_ids:
            where_sql = "WHERE contract_id IN (SELECT contract_id FROM app.contract_organizations WHERE organization_id = ANY(%s))"
            params = (organization_ids,)
        else:
            where_sql = "WHERE 1 = 0"
    with db() as conn:
        with conn.cursor() as cur:
            cur.execute(
                f"""
                SELECT *
                FROM app.pending_items
                {where_sql}
                ORDER BY
                    CASE status WHEN 'open' THEN 0 WHEN 'progress' THEN 1 WHEN 'done' THEN 2 ELSE 3 END,
                    due_date NULLS LAST,
                    id DESC
                """,
                params,
            )
            return [row_to_client(row) for row in cur.fetchall()]

def attach_company_contacts(cur, companies: list[dict[str, Any]]) -> None:
    company_ids = [parse_int(company.get("id")) for company in companies if parse_int(company.get("id"))]
    if not company_ids:
        return
    placeholders = ", ".join(["%s"] * len(company_ids))
    cur.execute(
        f"""
        SELECT *
        FROM app.company_contacts
        WHERE company_id IN ({placeholders})
        ORDER BY is_primary DESC, id
        """,
        tuple(company_ids),
    )
    contacts_by_company: dict[str, list[dict[str, Any]]] = {}
    for row in cur.fetchall():
        contact = row_to_client(row)
        contacts_by_company.setdefault(contact["companyId"], []).append(contact)
    for company in companies:
        contacts = contacts_by_company.get(company["id"], [])
        company["contacts"] = contacts
        primary = next((contact for contact in contacts if contact.get("isPrimary")), contacts[0] if contacts else None)
        if primary:
            company["manager"] = primary.get("name") or company.get("manager") or ""
            company["phone"] = primary.get("phone") or company.get("phone") or ""
            company["email"] = primary.get("email") or company.get("email") or ""


def fetch_one(table: str, item_id: int) -> dict[str, Any]:
    with db() as conn:
        with conn.cursor() as cur:
            cur.execute(f"SELECT * FROM {table} WHERE id = %s", (item_id,))
            row = cur.fetchone()
            if not row:
                raise HTTPException(status_code=404, detail="not found")
            return row_to_client(row)


def insert_row(table: str, values: dict[str, Any]) -> dict[str, Any]:
    columns = list(values.keys())
    placeholders = ", ".join(["%s"] * len(columns))
    sql = f"INSERT INTO {table} ({', '.join(columns)}) VALUES ({placeholders}) RETURNING *"
    with db() as conn:
        with conn.cursor() as cur:
            cur.execute(sql, tuple(values.values()))
            return row_to_client(cur.fetchone())


def update_row(table: str, item_id: int, values: dict[str, Any]) -> dict[str, Any]:
    if not values:
        return fetch_one(table, item_id)
    assignments = ", ".join([f"{column} = %s" for column in values])
    sql = f"UPDATE {table} SET {assignments} WHERE id = %s RETURNING *"
    with db() as conn:
        with conn.cursor() as cur:
            cur.execute(sql, (*values.values(), item_id))
            row = cur.fetchone()
            if not row:
                raise HTTPException(status_code=404, detail="not found")
            return row_to_client(row)


def delete_row(table: str, item_id: int) -> dict[str, bool]:
    with db() as conn:
        with conn.cursor() as cur:
            cur.execute(f"DELETE FROM {table} WHERE id = %s", (item_id,))
            if cur.rowcount == 0:
                raise HTTPException(status_code=404, detail="not found")
            return {"ok": True}


def has_column(table: str, column: str) -> bool:
    with db() as conn:
        with conn.cursor() as cur:
            cur.execute(
                """
                SELECT 1
                FROM information_schema.columns
                WHERE table_schema = 'public' AND table_name = %s AND column_name = %s
                """,
                (table, column),
            )
            return cur.fetchone() is not None


def parse_int(value: Any, default: int = 0) -> int:
    try:
        return int(value or default)
    except ValueError:
        return default


def parse_bool(value: Any) -> bool:
    if isinstance(value, bool):
        return value
    return str(value).lower() in {"1", "true", "yes", "on", "y"}


def normalize_text(text: str) -> str:
    return "\n".join(line.strip() for line in text.splitlines() if line.strip())


def ocr_image(path: Path) -> str:
    with Image.open(path) as image:
        return pytesseract.image_to_string(image, lang="kor+eng", config="--psm 6")


def ocr_pdf(path: Path) -> str:
    texts: list[str] = []
    document = fitz.open(path)
    try:
        for page_index in range(min(document.page_count, 2)):
            page = document.load_page(page_index)
            pixmap = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
            with Image.frombytes("RGB", [pixmap.width, pixmap.height], pixmap.samples) as image:
                texts.append(pytesseract.image_to_string(image, lang="kor+eng", config="--psm 6"))
    finally:
        document.close()
    return "\n".join(texts)


def extract_after_label(lines: list[str], labels: list[str]) -> str:
    for index, line in enumerate(lines):
        compact = re.sub(r"\s+", "", line)
        for label in labels:
            label_compact = re.sub(r"\s+", "", label)
            if label_compact in compact:
                label_pattern = r"\s*".join(re.escape(char) for char in label)
                after = re.sub(
                    rf"^.*?{label_pattern}\s*[:：\-\|ㆍ·]?\s*",
                    "",
                    line,
                    count=1,
                    flags=re.IGNORECASE,
                ).strip(":-|ㆍ·()[]{} ")
                if after and len(after) > 1:
                    return clean_ocr_field_value(after, labels)
                if index + 1 < len(lines):
                    return clean_ocr_field_value(lines[index + 1], labels)
    return ""


def clean_ocr_field_value(value: str, labels: list[str]) -> str:
    cleaned = re.sub(r"\s+", " ", value).strip(":-|ㆍ·()[]{} \t")
    for _ in range(3):
        before = cleaned
        for label in labels:
            escaped = re.escape(label)
            cleaned = re.sub(
                rf"^\s*[\(\[\{{]?\s*{escaped}\s*[\)\]\}}]?\s*[:：\-\|ㆍ·]?\s*",
                "",
                cleaned,
                flags=re.IGNORECASE,
            ).strip(":-|ㆍ·()[]{} \t")
        if cleaned == before:
            break
    return cleaned


def parse_company_ocr(text: str) -> dict[str, str]:
    cleaned = normalize_text(text)
    lines = cleaned.splitlines()
    joined = " ".join(lines)
    name_labels = ["단체명", "상호", "법인명", "업체명", "회사명", "상호명"]
    representative_labels = ["대표자", "성명", "대표"]
    manager_labels = ["담당자"]
    address_labels = ["사업장 소재지", "소재지", "주소", "본점소재지"]
    business_no = ""
    match = re.search(r"\b\d{3}[-\s]?\d{2}[-\s]?\d{5}\b", joined)
    if match:
        digits = re.sub(r"\D", "", match.group(0))
        business_no = f"{digits[:3]}-{digits[3:5]}-{digits[5:]}"

    email = ""
    email_match = re.search(r"[\w.+-]+@[\w.-]+\.[A-Za-z]{2,}", joined)
    if email_match:
        email = email_match.group(0)

    phone = ""
    phone_match = re.search(r"(?:0\d{1,2}[-\s]?)?\d{3,4}[-\s]?\d{4}", joined)
    if phone_match:
        phone = re.sub(r"\s+", "-", phone_match.group(0)).replace("--", "-")

    name = extract_after_label(lines, name_labels)
    representative = extract_after_label(lines, representative_labels)
    manager = extract_after_label(lines, manager_labels)
    address = extract_after_label(lines, address_labels)

    return {
        "name": clean_ocr_field_value(name, name_labels),
        "businessNo": business_no,
        "representative": clean_ocr_field_value(representative, representative_labels),
        "address": clean_ocr_field_value(address, address_labels),
        "manager": clean_ocr_field_value(manager, manager_labels),
        "phone": phone,
        "email": email,
        "rawText": cleaned[:5000],
    }


app = FastAPI(title="Maintenance Contract Manager")
SESSION_COOKIE = "maintenance_session"
SESSION_MAX_AGE = 60 * 60 * 2


def password_hash(password: str) -> str:
    salt = secrets_token(16)
    iterations = 120_000
    digest = hashlib.pbkdf2_hmac("sha256", password.encode("utf-8"), salt.encode("utf-8"), iterations)
    return f"pbkdf2_sha256${iterations}${salt}${base64.b64encode(digest).decode('ascii')}"


def verify_password(password: str, stored: str) -> bool:
    try:
        algorithm, iterations_text, salt, digest_text = stored.split("$", 3)
        if algorithm != "pbkdf2_sha256":
            return False
        digest = hashlib.pbkdf2_hmac("sha256", password.encode("utf-8"), salt.encode("utf-8"), int(iterations_text))
        return hmac.compare_digest(base64.b64encode(digest).decode("ascii"), digest_text)
    except (ValueError, TypeError):
        return False


def validate_signup_password(password: str) -> None:
    if len(password) < 6 or len(password) >= 12:
        raise HTTPException(status_code=400, detail="비밀번호는 6자리 이상 12자리 미만이어야 합니다.")
    if not re.search(r"[A-Za-z]", password):
        raise HTTPException(status_code=400, detail="비밀번호에는 영문이 포함되어야 합니다.")
    if not re.search(r"[^A-Za-z0-9]", password):
        raise HTTPException(status_code=400, detail="비밀번호에는 특수문자가 포함되어야 합니다.")
    if re.search(r"\s", password):
        raise HTTPException(status_code=400, detail="비밀번호에는 공백을 사용할 수 없습니다.")


def seed_organizations(cur) -> None:
    cur.execute("SELECT 1 FROM app.organizations WHERE is_active = TRUE LIMIT 1")
    if cur.fetchone():
        return
    org_ids: dict[str, int] = {}
    for index, (name, parent_name) in enumerate(DEFAULT_ORGANIZATIONS, start=1):
        parent_id = org_ids.get(parent_name or "")
        cur.execute(
            """
            INSERT INTO app.organizations (name, parent_id, sort_order, updated_at)
            VALUES (%s, %s, %s, now())
            RETURNING id
            """,
            (name, parent_id, index),
        )
        org_ids[name] = int(cur.fetchone()["id"])


def secrets_token(length: int = 32) -> str:
    return base64.urlsafe_b64encode(os.urandom(length)).decode("ascii").rstrip("=")


def sign_session(user_id: int, username: str, role: str) -> str:
    expires = int(time()) + SESSION_MAX_AGE
    payload = f"{user_id}|{username}|{role}|{expires}|{secrets_token(12)}"
    signature = hmac.new(CONFIG["SESSION_SECRET"].encode("utf-8"), payload.encode("utf-8"), hashlib.sha256).hexdigest()
    return base64.urlsafe_b64encode(f"{payload}|{signature}".encode("utf-8")).decode("ascii")


def read_session(token: str | None) -> dict[str, Any] | None:
    if not token:
        return None
    try:
        decoded = base64.urlsafe_b64decode(token.encode("ascii")).decode("utf-8")
        user_id, username, role, expires_text, nonce, signature = decoded.rsplit("|", 5)
        payload = "|".join([user_id, username, role, expires_text, nonce])
        expected = hmac.new(CONFIG["SESSION_SECRET"].encode("utf-8"), payload.encode("utf-8"), hashlib.sha256).hexdigest()
        if not hmac.compare_digest(signature, expected) or int(expires_text) < int(time()):
            return None
        return {"id": int(user_id), "username": username, "role": role}
    except (ValueError, TypeError, UnicodeDecodeError):
        return None


def init_auth_schema() -> None:
    with db() as conn:
        with conn.cursor() as cur:
            cur.execute("CREATE SCHEMA IF NOT EXISTS app AUTHORIZATION maintenance_user")
            cur.execute(
                """
                CREATE TABLE IF NOT EXISTS app.organizations (
                    id BIGSERIAL PRIMARY KEY,
                    name TEXT NOT NULL,
                    parent_id BIGINT REFERENCES app.organizations(id) ON DELETE SET NULL,
                    sort_order INTEGER NOT NULL DEFAULT 0,
                    is_active BOOLEAN NOT NULL DEFAULT TRUE,
                    created_at TIMESTAMP NOT NULL DEFAULT now(),
                    updated_at TIMESTAMP NOT NULL DEFAULT now()
                )
                """
            )
            cur.execute("ALTER TABLE app.organizations DROP CONSTRAINT IF EXISTS organizations_name_key")
            cur.execute(
                """
                CREATE TABLE IF NOT EXISTS app.app_users (
                    id BIGSERIAL PRIMARY KEY,
                    username TEXT UNIQUE NOT NULL,
                    password_hash TEXT NOT NULL,
                    display_name TEXT,
                    role TEXT NOT NULL DEFAULT 'staff',
                    is_active BOOLEAN NOT NULL DEFAULT TRUE,
                    created_at TIMESTAMP NOT NULL DEFAULT now(),
                    updated_at TIMESTAMP NOT NULL DEFAULT now()
                )
                """
            )
            cur.execute("ALTER TABLE app.app_users ADD COLUMN IF NOT EXISTS organization_id BIGINT REFERENCES app.organizations(id) ON DELETE SET NULL")
            cur.execute("ALTER TABLE app.app_users ADD COLUMN IF NOT EXISTS email TEXT")
            cur.execute(
                """
                CREATE TABLE IF NOT EXISTS app.employee_master (
                    id BIGSERIAL PRIMARY KEY,
                    employee_no TEXT NOT NULL,
                    name TEXT NOT NULL,
                    department_code TEXT,
                    position_code TEXT,
                    organization_id BIGINT REFERENCES app.organizations(id) ON DELETE SET NULL,
                    position TEXT,
                    email TEXT,
                    phone TEXT,
                    employment_status TEXT NOT NULL DEFAULT 'active',
                    is_active BOOLEAN NOT NULL DEFAULT TRUE,
                    joined_at DATE,
                    resigned_at DATE,
                    memo TEXT,
                    created_at TIMESTAMP NOT NULL DEFAULT now(),
                    updated_at TIMESTAMP NOT NULL DEFAULT now()
                )
                """
            )
            cur.execute(
                """
                CREATE TABLE IF NOT EXISTS app.employee_departments (
                    code TEXT PRIMARY KEY,
                    parent_name TEXT,
                    name TEXT NOT NULL,
                    raw_name TEXT,
                    is_active BOOLEAN NOT NULL DEFAULT TRUE,
                    created_at TIMESTAMP NOT NULL DEFAULT now(),
                    updated_at TIMESTAMP NOT NULL DEFAULT now()
                )
                """
            )
            cur.execute(
                """
                CREATE TABLE IF NOT EXISTS app.employee_positions (
                    code TEXT PRIMARY KEY,
                    name TEXT NOT NULL,
                    is_active BOOLEAN NOT NULL DEFAULT TRUE,
                    created_at TIMESTAMP NOT NULL DEFAULT now(),
                    updated_at TIMESTAMP NOT NULL DEFAULT now()
                )
                """
            )
            cur.execute("ALTER TABLE app.app_users ADD COLUMN IF NOT EXISTS employee_id BIGINT REFERENCES app.employee_master(id) ON DELETE SET NULL")
            cur.execute("ALTER TABLE app.employee_master DROP CONSTRAINT IF EXISTS employee_master_employee_no_key")
            cur.execute("ALTER TABLE app.employee_master ADD COLUMN IF NOT EXISTS department_code TEXT")
            cur.execute("ALTER TABLE app.employee_master ADD COLUMN IF NOT EXISTS position_code TEXT")
            cur.execute(
                """
                CREATE TABLE IF NOT EXISTS app.company_profiles (
                    company_id INTEGER PRIMARY KEY REFERENCES public.companies(id) ON DELETE CASCADE,
                    representative TEXT,
                    updated_at TIMESTAMP NOT NULL DEFAULT now()
                )
                """
            )
            cur.execute(
                """
                CREATE TABLE IF NOT EXISTS app.company_organizations (
                    company_id INTEGER PRIMARY KEY REFERENCES public.companies(id) ON DELETE CASCADE,
                    organization_id BIGINT REFERENCES app.organizations(id) ON DELETE SET NULL,
                    created_at TIMESTAMP NOT NULL DEFAULT now(),
                    updated_at TIMESTAMP NOT NULL DEFAULT now()
                )
                """
            )
            cur.execute(
                """
                CREATE TABLE IF NOT EXISTS app.contract_organizations (
                    contract_id INTEGER PRIMARY KEY REFERENCES public.contracts(id) ON DELETE CASCADE,
                    organization_id BIGINT REFERENCES app.organizations(id) ON DELETE SET NULL,
                    created_at TIMESTAMP NOT NULL DEFAULT now(),
                    updated_at TIMESTAMP NOT NULL DEFAULT now()
                )
                """
            )
            cur.execute(
                """
                CREATE TABLE IF NOT EXISTS app.organization_contract_links (
                    id BIGSERIAL PRIMARY KEY,
                    organization_a_id BIGINT NOT NULL REFERENCES app.organizations(id) ON DELETE CASCADE,
                    organization_b_id BIGINT NOT NULL REFERENCES app.organizations(id) ON DELETE CASCADE,
                    is_active BOOLEAN NOT NULL DEFAULT TRUE,
                    created_at TIMESTAMP NOT NULL DEFAULT now(),
                    updated_at TIMESTAMP NOT NULL DEFAULT now(),
                    CONSTRAINT organization_contract_links_distinct CHECK (organization_a_id < organization_b_id),
                    CONSTRAINT organization_contract_links_pair_unique UNIQUE (organization_a_id, organization_b_id)
                )
                """
            )
            cur.execute(
                """
                CREATE TABLE IF NOT EXISTS app.contract_details (
                    contract_id INTEGER PRIMARY KEY REFERENCES public.contracts(id) ON DELETE CASCADE,
                    department TEXT,
                    manager TEXT,
                    manager_user_id BIGINT REFERENCES app.app_users(id) ON DELETE SET NULL,
                    created_at TIMESTAMP NOT NULL DEFAULT now(),
                    updated_at TIMESTAMP NOT NULL DEFAULT now()
                )
                """
            )
            cur.execute("ALTER TABLE app.contract_details ADD COLUMN IF NOT EXISTS manager_user_id BIGINT REFERENCES app.app_users(id) ON DELETE SET NULL")
            cur.execute(
                """
                CREATE TABLE IF NOT EXISTS app.document_organizations (
                    document_id INTEGER PRIMARY KEY REFERENCES public.documents(id) ON DELETE CASCADE,
                    organization_id BIGINT REFERENCES app.organizations(id) ON DELETE SET NULL,
                    created_at TIMESTAMP NOT NULL DEFAULT now(),
                    updated_at TIMESTAMP NOT NULL DEFAULT now()
                )
                """
            )
            cur.execute(
                """
                CREATE TABLE IF NOT EXISTS app.company_contacts (
                    id BIGSERIAL PRIMARY KEY,
                    company_id INTEGER NOT NULL REFERENCES public.companies(id) ON DELETE CASCADE,
                    name TEXT NOT NULL,
                    position TEXT,
                    phone TEXT,
                    email TEXT,
                    memo TEXT,
                    is_primary BOOLEAN NOT NULL DEFAULT FALSE,
                    created_at TIMESTAMP NOT NULL DEFAULT now(),
                    updated_at TIMESTAMP NOT NULL DEFAULT now()
                )
                """
            )
            cur.execute(
                """
                CREATE TABLE IF NOT EXISTS app.pending_items (
                    id BIGSERIAL PRIMARY KEY,
                    company_id INTEGER REFERENCES public.companies(id) ON DELETE SET NULL,
                    contract_id INTEGER REFERENCES public.contracts(id) ON DELETE SET NULL,
                    title TEXT NOT NULL,
                    memo TEXT,
                    due_date DATE,
                    status TEXT NOT NULL DEFAULT 'open',
                    created_at TIMESTAMP NOT NULL DEFAULT now(),
                    updated_at TIMESTAMP NOT NULL DEFAULT now(),
                    completed_at TIMESTAMP
                )
                """
            )
            cur.execute(
                """
                CREATE TABLE IF NOT EXISTS app.audit_logs (
                    id BIGSERIAL PRIMARY KEY,
                    user_id BIGINT REFERENCES app.app_users(id) ON DELETE SET NULL,
                    username TEXT,
                    action TEXT NOT NULL,
                    target_table TEXT NOT NULL,
                    target_id TEXT,
                    detail TEXT,
                    created_at TIMESTAMP NOT NULL DEFAULT now()
                )
                """
            )
            cur.execute(
                """
                CREATE TABLE IF NOT EXISTS app.email_notification_logs (
                    id BIGSERIAL PRIMARY KEY,
                    contract_id INTEGER REFERENCES public.contracts(id) ON DELETE CASCADE,
                    user_id BIGINT REFERENCES app.app_users(id) ON DELETE SET NULL,
                    recipient_email TEXT NOT NULL,
                    alert_days INTEGER NOT NULL,
                    status TEXT NOT NULL DEFAULT 'sent',
                    detail TEXT,
                    sent_at TIMESTAMP NOT NULL DEFAULT now()
                )
                """
            )
            cur.execute("CREATE INDEX IF NOT EXISTS idx_pending_items_status_due ON app.pending_items (status, due_date)")
            cur.execute("CREATE INDEX IF NOT EXISTS idx_pending_items_company ON app.pending_items (company_id)")
            cur.execute("CREATE INDEX IF NOT EXISTS idx_pending_items_contract ON app.pending_items (contract_id)")
            cur.execute("CREATE INDEX IF NOT EXISTS idx_app_users_organization ON app.app_users (organization_id)")
            cur.execute("CREATE INDEX IF NOT EXISTS idx_app_users_employee ON app.app_users (employee_id)")
            cur.execute("CREATE UNIQUE INDEX IF NOT EXISTS idx_app_users_employee_unique ON app.app_users (employee_id) WHERE employee_id IS NOT NULL")
            cur.execute("CREATE INDEX IF NOT EXISTS idx_employee_master_organization ON app.employee_master (organization_id)")
            cur.execute("CREATE INDEX IF NOT EXISTS idx_employee_master_status ON app.employee_master (employment_status, is_active)")
            cur.execute("CREATE INDEX IF NOT EXISTS idx_employee_master_department_code ON app.employee_master (department_code)")
            cur.execute("CREATE INDEX IF NOT EXISTS idx_employee_master_position_code ON app.employee_master (position_code)")
            cur.execute("CREATE UNIQUE INDEX IF NOT EXISTS idx_employee_master_no_name_unique ON app.employee_master (employee_no, name)")
            cur.execute("CREATE INDEX IF NOT EXISTS idx_company_organizations_org ON app.company_organizations (organization_id)")
            cur.execute("CREATE INDEX IF NOT EXISTS idx_contract_organizations_org ON app.contract_organizations (organization_id)")
            cur.execute("CREATE INDEX IF NOT EXISTS idx_org_contract_links_a ON app.organization_contract_links (organization_a_id) WHERE is_active = TRUE")
            cur.execute("CREATE INDEX IF NOT EXISTS idx_org_contract_links_b ON app.organization_contract_links (organization_b_id) WHERE is_active = TRUE")
            cur.execute("CREATE INDEX IF NOT EXISTS idx_contract_details_contract ON app.contract_details (contract_id)")
            cur.execute("CREATE INDEX IF NOT EXISTS idx_document_organizations_org ON app.document_organizations (organization_id)")
            cur.execute("CREATE INDEX IF NOT EXISTS idx_company_contacts_company ON app.company_contacts (company_id)")
            cur.execute("CREATE UNIQUE INDEX IF NOT EXISTS idx_company_contacts_primary ON app.company_contacts (company_id) WHERE is_primary")
            cur.execute("CREATE INDEX IF NOT EXISTS idx_audit_logs_created_at ON app.audit_logs (created_at DESC)")
            cur.execute("CREATE INDEX IF NOT EXISTS idx_audit_logs_target ON app.audit_logs (target_table, target_id)")
            cur.execute("CREATE INDEX IF NOT EXISTS idx_audit_logs_user ON app.audit_logs (user_id)")
            cur.execute("CREATE UNIQUE INDEX IF NOT EXISTS idx_organizations_active_parent_name ON app.organizations (COALESCE(parent_id, 0), name) WHERE is_active = TRUE")
            cur.execute("CREATE INDEX IF NOT EXISTS idx_app_users_email ON app.app_users (email)")
            cur.execute("CREATE UNIQUE INDEX IF NOT EXISTS idx_email_logs_once ON app.email_notification_logs (contract_id, user_id, alert_days)")
            cur.execute(
                """
                UPDATE app.contract_details cd
                SET manager_user_id = u.id,
                    updated_at = now()
                FROM app.app_users u
                WHERE cd.manager_user_id IS NULL
                  AND NULLIF(cd.manager, '') IS NOT NULL
                  AND lower(trim(cd.manager)) IN (
                      lower(trim(COALESCE(NULLIF(u.display_name, ''), u.username))),
                      lower(trim(u.username))
                  )
                """
            )
            cur.execute(
                """
                INSERT INTO app.company_contacts (company_id, name, phone, email, is_primary)
                SELECT c.id, COALESCE(NULLIF(c.manager, ''), '담당자'), c.phone, c.email, TRUE
                FROM public.companies c
                WHERE NOT EXISTS (
                    SELECT 1 FROM app.company_contacts cc WHERE cc.company_id = c.id
                )
                AND (
                    NULLIF(c.manager, '') IS NOT NULL
                    OR NULLIF(c.phone, '') IS NOT NULL
                    OR NULLIF(c.email, '') IS NOT NULL
                )
                """
            )
            seed_organizations(cur)
            cur.execute(
                """
                INSERT INTO app.organization_contract_links
                    (organization_a_id, organization_b_id, is_active, updated_at)
                SELECT LEAST(id, parent_id), GREATEST(id, parent_id), TRUE, now()
                FROM app.organizations
                WHERE parent_id IS NOT NULL AND is_active = TRUE
                ON CONFLICT (organization_a_id, organization_b_id) DO NOTHING
                """
            )
            cur.execute("UPDATE app.app_users SET role = 'team_member' WHERE role IN ('staff', 'viewer')")
            cur.execute("UPDATE app.app_users SET role = 'team_lead' WHERE role = 'leader'")
            cur.execute("SELECT id FROM app.organizations WHERE name = %s", ("경영지원실",))
            default_org = cur.fetchone()
            if default_org:
                cur.execute("UPDATE app.app_users SET organization_id = %s WHERE organization_id IS NULL AND role <> 'admin'", (default_org["id"],))
                cur.execute(
                    """
                    INSERT INTO app.contract_organizations (contract_id, organization_id, updated_at)
                    SELECT c.id, %s, now()
                    FROM public.contracts c
                    WHERE NOT EXISTS (
                        SELECT 1 FROM app.contract_organizations co WHERE co.contract_id = c.id
                    )
                    """,
                    (default_org["id"],),
                )
                cur.execute(
                    """
                    INSERT INTO app.document_organizations (document_id, organization_id, updated_at)
                    SELECT d.id, COALESCE(co.organization_id, %s), now()
                    FROM public.documents d
                    LEFT JOIN app.contract_organizations co ON co.contract_id = d.contract_id
                    WHERE NOT EXISTS (
                        SELECT 1 FROM app.document_organizations doo WHERE doo.document_id = d.id
                    )
                    """,
                    (default_org["id"],),
                )
            cur.execute("SELECT id FROM app.app_users WHERE username = %s", ("admin",))
            if cur.fetchone() is None:
                cur.execute(
                    """
                    INSERT INTO app.app_users (username, password_hash, display_name, role, is_active)
                    VALUES (%s, %s, %s, %s, TRUE)
                    """,
                    ("admin", password_hash("Dk@dmin2014"), "관리자", "admin"),
                )
            default_email = CONFIG.get("RECEIVER_EMAIL") or CONFIG.get("SENDER_EMAIL")
            if default_email:
                cur.execute(
                    """
                    UPDATE app.app_users
                    SET email = %s, updated_at = now()
                    WHERE username IN ('admin', 'jkoper')
                       OR display_name IN ('정광', '관리자')
                    """,
                    (default_email,),
                )


@app.on_event("startup")
def startup() -> None:
    init_auth_schema()
    start_mail_scheduler()


def client_user(row: dict[str, Any]) -> dict[str, Any]:
    return {
        "id": str(row["id"]),
        "username": row["username"],
        "employeeId": "" if row.get("employee_id") is None else str(row.get("employee_id")),
        "employeeNo": row.get("employee_no") or "",
        "displayName": row.get("display_name") or "",
        "email": row.get("email") or "",
        "role": row.get("role") or "staff",
        "positionName": row.get("position_name") or row.get("position") or "",
        "position": row.get("position_name") or row.get("position") or "",
        "organizationId": "" if row.get("organization_id") is None else str(row.get("organization_id")),
        "organizationName": row.get("organization_name") or "",
        "isActive": bool(row.get("is_active")),
        "createdAt": row.get("created_at").isoformat() if row.get("created_at") else "",
        "updatedAt": row.get("updated_at").isoformat() if row.get("updated_at") else "",
    }


def resolve_employee_organization_id(cur, employee: dict[str, Any]) -> int | None:
    if employee.get("organization_id"):
        return int(employee["organization_id"])
    cur.execute(
        """
        SELECT d.parent_name, d.name
        FROM app.employee_departments d
        WHERE d.code = %s
        """,
        (employee.get("department_code"),),
    )
    department = cur.fetchone() or {}
    names = department_organization_candidates(department.get("parent_name"), department.get("name"))
    for name in names:
        if not name:
            continue
        cur.execute("SELECT id FROM app.organizations WHERE is_active = TRUE AND name = %s LIMIT 1", (name,))
        row = cur.fetchone()
        if row:
            return int(row["id"])
    for name in names:
        key = normalize_organization_name(name)
        if not key:
            continue
        cur.execute(
            """
            SELECT id
            FROM app.organizations
            WHERE is_active = TRUE
              AND regexp_replace(lower(name), '[\\s\\-_·.]+', '', 'g') = %s
            ORDER BY id
            LIMIT 2
            """,
            (key,),
        )
        rows = cur.fetchall()
        if len(rows) == 1:
            return int(rows[0]["id"])
    return None


def fetch_organizations() -> list[dict[str, Any]]:
    with db() as conn:
        with conn.cursor() as cur:
            cur.execute("SELECT * FROM app.organizations WHERE is_active = TRUE ORDER BY sort_order, id")
            return [row_to_client(row) for row in cur.fetchall()]


def fetch_organization_contract_links() -> list[dict[str, Any]]:
    with db() as conn:
        with conn.cursor() as cur:
            cur.execute(
                """
                SELECT l.id, l.organization_a_id, l.organization_b_id,
                       a.name AS organization_a_name, b.name AS organization_b_name
                FROM app.organization_contract_links l
                JOIN app.organizations a ON a.id = l.organization_a_id AND a.is_active = TRUE
                JOIN app.organizations b ON b.id = l.organization_b_id AND b.is_active = TRUE
                WHERE l.is_active = TRUE
                  AND a.parent_id IS NOT NULL
                  AND b.parent_id IS NOT NULL
                ORDER BY a.name, b.name, l.id
                """
            )
            return [
                {
                    "id": str(row["id"]),
                    "organizationAId": str(row["organization_a_id"]),
                    "organizationBId": str(row["organization_b_id"]),
                    "organizationAName": row.get("organization_a_name") or "",
                    "organizationBName": row.get("organization_b_name") or "",
                }
                for row in cur.fetchall()
            ]


def contract_organization_source_id(user: dict[str, Any], cur) -> int | None:
    organization_id = parse_int(user.get("organizationId"))
    if not organization_id:
        return None
    cur.execute(
        """
        SELECT child.id, child.parent_id, parent.parent_id AS grand_parent_id
        FROM app.organizations child
        LEFT JOIN app.organizations parent
          ON parent.id = child.parent_id AND parent.is_active = TRUE
        WHERE child.id = %s AND child.is_active = TRUE
        """,
        (organization_id,),
    )
    row = cur.fetchone()
    if not row:
        return None
    if (
        user.get("role") == "team_member"
        and row.get("parent_id")
        and row.get("grand_parent_id") is not None
    ):
        return int(row["parent_id"])
    return organization_id


def get_allowed_contract_organization_ids(user: dict[str, Any]) -> list[str]:
    role = user.get("role")
    with db() as conn:
        with conn.cursor() as cur:
            if role in FULL_ACCESS_ROLES:
                cur.execute(
                    """
                    SELECT id
                    FROM app.organizations
                    WHERE is_active = TRUE AND parent_id IS NOT NULL
                    ORDER BY sort_order, id
                    """
                )
                return [str(row["id"]) for row in cur.fetchall()]

            source_id = contract_organization_source_id(user, cur)
            if not source_id:
                return []
            cur.execute(
                """
                SELECT 1
                FROM app.organizations
                WHERE id = %s AND is_active = TRUE AND parent_id IS NOT NULL
                """,
                (source_id,),
            )
            if not cur.fetchone():
                return []
            cur.execute(
                """
                SELECT l.organization_a_id, l.organization_b_id
                FROM app.organization_contract_links l
                JOIN app.organizations a ON a.id = l.organization_a_id AND a.is_active = TRUE
                JOIN app.organizations b ON b.id = l.organization_b_id AND b.is_active = TRUE
                WHERE l.is_active = TRUE
                  AND a.parent_id IS NOT NULL
                  AND b.parent_id IS NOT NULL
                  AND %s IN (l.organization_a_id, l.organization_b_id)
                """,
                (source_id,),
            )
            resolved_ids = resolve_allowed_organization_ids(source_id, cur.fetchall())
            ordered_ids = [source_id, *sorted(value for value in resolved_ids if value != source_id)]
            return [str(value) for value in ordered_ids]


def replace_organization_contract_links(source_id: int, target_ids: list[Any]) -> None:
    normalized_targets = {parse_int(value) for value in target_ids}
    normalized_targets.discard(0)
    normalized_targets.discard(source_id)
    with db() as conn:
        with conn.cursor() as cur:
            cur.execute(
                """
                SELECT 1
                FROM app.organizations
                WHERE id = %s AND is_active = TRUE AND parent_id IS NOT NULL
                """,
                (source_id,),
            )
            if not cur.fetchone():
                raise HTTPException(status_code=400, detail="연계 기준 조직을 찾을 수 없습니다.")

            valid_targets: set[int] = set()
            if normalized_targets:
                cur.execute(
                    """
                    SELECT id
                    FROM app.organizations
                    WHERE id = ANY(%s) AND is_active = TRUE AND parent_id IS NOT NULL
                    """,
                    (sorted(normalized_targets),),
                )
                valid_targets = {int(row["id"]) for row in cur.fetchall()}
                if valid_targets != normalized_targets:
                    raise HTTPException(status_code=400, detail="연계할 수 없는 조직이 포함되어 있습니다.")

            cur.execute(
                """
                DELETE FROM app.organization_contract_links
                WHERE organization_a_id = %s OR organization_b_id = %s
                """,
                (source_id, source_id),
            )
            for target_id in sorted(valid_targets):
                first_id, second_id = normalize_organization_pair(source_id, target_id)
                cur.execute(
                    """
                    INSERT INTO app.organization_contract_links
                        (organization_a_id, organization_b_id, is_active, updated_at)
                    VALUES (%s, %s, TRUE, now())
                    ON CONFLICT (organization_a_id, organization_b_id)
                    DO UPDATE SET is_active = TRUE, updated_at = now()
                    """,
                    (first_id, second_id),
                )


def fetch_users_for_admin() -> list[dict[str, Any]]:
    with db() as conn:
        with conn.cursor() as cur:
            cur.execute(
                """
                SELECT u.*, o.name AS organization_name, e.employee_no,
                       p.name AS position_name, e.position
                FROM app.app_users u
                LEFT JOIN app.organizations o ON o.id = u.organization_id
                LEFT JOIN app.employee_master e ON e.id = u.employee_id
                LEFT JOIN app.employee_positions p ON p.code = e.position_code
                ORDER BY u.id
                """
            )
            return [client_user(row) for row in cur.fetchall()]


def client_employee(row: dict[str, Any]) -> dict[str, Any]:
    return {
        "id": str(row["id"]),
        "employeeNo": row.get("employee_no") or "",
        "name": row.get("name") or "",
        "departmentCode": row.get("department_code") or "",
        "departmentParentName": row.get("department_parent_name") or "",
        "departmentName": row.get("department_name") or "",
        "positionCode": row.get("position_code") or "",
        "positionName": row.get("position_name") or row.get("position") or "",
        "organizationId": "" if row.get("organization_id") is None else str(row.get("organization_id")),
        "organizationName": row.get("organization_name") or "",
        "position": row.get("position_name") or row.get("position") or "",
        "email": row.get("email") or "",
        "phone": row.get("phone") or "",
        "employmentStatus": row.get("employment_status") or "active",
        "isActive": bool(row.get("is_active")),
        "joinedAt": row.get("joined_at").isoformat() if row.get("joined_at") else "",
        "resignedAt": row.get("resigned_at").isoformat() if row.get("resigned_at") else "",
        "memo": row.get("memo") or "",
        "createdAt": row.get("created_at").isoformat() if row.get("created_at") else "",
        "updatedAt": row.get("updated_at").isoformat() if row.get("updated_at") else "",
    }


def fetch_employees() -> list[dict[str, Any]]:
    with db() as conn:
        with conn.cursor() as cur:
            cur.execute(
                """
                SELECT e.*, o.name AS organization_name,
                       d.parent_name AS department_parent_name, d.name AS department_name,
                       p.name AS position_name
                FROM app.employee_master e
                LEFT JOIN app.organizations o ON o.id = e.organization_id
                LEFT JOIN app.employee_departments d ON d.code = e.department_code
                LEFT JOIN app.employee_positions p ON p.code = e.position_code
                ORDER BY e.is_active DESC, e.employee_no, e.id
                """
            )
            return [client_employee(row) for row in cur.fetchall()]


def fetch_selectable_users(user: dict[str, Any]) -> list[dict[str, Any]]:
    where_sql = "WHERE u.is_active = TRUE"
    params: tuple[Any, ...] = ()
    role = user.get("role")
    if role not in FULL_ACCESS_ROLES:
        organization_id = parse_int(user.get("organizationId"))
        if role == "team_lead" and organization_id:
            where_sql += " AND (u.id = %s OR (u.organization_id = %s AND u.role = 'team_member'))"
            params = (parse_int(user.get("id")), organization_id)
        else:
            where_sql += " AND u.id = %s"
            params = (parse_int(user.get("id")),)
    with db() as conn:
        with conn.cursor() as cur:
            cur.execute(
                f"""
                SELECT u.*, o.name AS organization_name, e.employee_no,
                       p.name AS position_name, e.position
                FROM app.app_users u
                LEFT JOIN app.organizations o ON o.id = u.organization_id
                LEFT JOIN app.employee_master e ON e.id = u.employee_id
                LEFT JOIN app.employee_positions p ON p.code = e.position_code
                {where_sql}
                  AND (u.employee_id IS NULL OR (e.is_active = TRUE AND e.employment_status = 'active'))
                ORDER BY o.name NULLS LAST, u.display_name NULLS LAST, u.username
                """,
                params,
            )
            return [client_user(row) for row in cur.fetchall()]


def current_user_full(request: Request) -> dict[str, Any]:
    user = getattr(request.state, "user", None)
    if not user:
        raise HTTPException(status_code=401, detail="login required")
    with db() as conn:
        with conn.cursor() as cur:
            cur.execute(
                """
                SELECT u.*, o.name AS organization_name, e.employee_no,
                       p.name AS position_name, e.position
                FROM app.app_users u
                LEFT JOIN app.organizations o ON o.id = u.organization_id
                LEFT JOIN app.employee_master e ON e.id = u.employee_id
                LEFT JOIN app.employee_positions p ON p.code = e.position_code
                WHERE u.id = %s AND u.is_active = TRUE
                  AND (u.employee_id IS NULL OR (e.is_active = TRUE AND e.employment_status = 'active'))
                """,
                (user["id"],),
            )
            row = cur.fetchone()
    if not row:
        raise HTTPException(status_code=401, detail="login required")
    return client_user(row)


def ensure_permission(request: Request, permission: str) -> dict[str, Any]:
    user = current_user_full(request)
    allowed_permissions = ROLE_PERMISSIONS.get(user.get("role") or "", set())
    if permission not in allowed_permissions:
        raise HTTPException(status_code=403, detail="권한이 없습니다.")
    return user


def ensure_write_permission(request: Request, permission: str) -> dict[str, Any]:
    return ensure_permission(request, permission)


def ensure_can_delete(request: Request) -> None:
    user = current_user_full(request)
    if user and user.get("role") == "team_member":
        raise HTTPException(status_code=403, detail="팀원 권한은 삭제할 수 없습니다.")


def ensure_company_update_permission(request: Request) -> dict[str, Any]:
    return ensure_permission(request, "menu.company.update")


CONTRACT_ASSIGNMENT_ERROR = "해당 부서 및 담당자 등록 권한이 없습니다"


def ensure_contract_create_permission(request: Request) -> dict[str, Any]:
    return ensure_permission(request, "menu.contract.create")


def ensure_contract_update_permission(request: Request) -> dict[str, Any]:
    return ensure_permission(request, "menu.contract.update")


def ensure_contract_delete_permission(item_id: int, request: Request) -> None:
    ensure_permission(request, "menu.contract.delete")


def ensure_admin_delete(request: Request) -> None:
    user = current_user_full(request)
    if user.get("role") not in FULL_ACCESS_ROLES:
        raise HTTPException(status_code=403, detail="삭제 권한이 없습니다.")


def ensure_excel_import_admin(request: Request) -> dict[str, Any]:
    user = current_user_full(request)
    if user.get("role") != "admin":
        raise HTTPException(status_code=403, detail="Excel 일괄 등록은 관리자만 사용할 수 있습니다.")
    return user


EXCEL_TEMPLATES = {
    "companies": {
        "filename": "업체_일괄등록_양식.xlsx",
        "permission": "menu.company.view",
        "headers": ["업체명", "사업자번호", "대표자", "주소", "담당자명", "담당자연락처", "담당자이메일", "비고"],
        "guides": ["업체명은 필수입니다.", "기존 업체명 또는 사업자번호와 중복되면 등록되지 않습니다."],
    },
    "contracts": {
        "filename": "계약_일괄등록_양식.xlsx",
        "permission": "menu.contract.view",
        "headers": ["업체사업자번호", "업체명", "계약명", "담당부서", "담당자아이디", "시작일", "종료일", "금액", "지급방법", "상태", "갱신일", "알림일수", "자동알림", "비고"],
        "guides": ["날짜는 YYYY-MM-DD 형식으로 입력합니다.", "지급방법: 일시불, 3분할, 2분할, 월", "상태: 진행, 검토, 만료, 종료", "자동알림: 예 또는 아니오"],
    },
    "worklogs": {
        "filename": "유지보수이력_일괄등록_양식.xlsx",
        "permission": "menu.worklog.view",
        "headers": ["업체사업자번호", "업체명", "계약명", "작업일", "작업구분", "처리상태", "처리자", "제목", "작업내용"],
        "guides": ["작업일은 YYYY-MM-DD 형식으로 입력합니다.", "작업구분: 방문 점검, 원격지원, 장애 처리, 부품 교체, 정기점검", "처리상태: 접수, 처리중, 완료, 보류", "업체와 계약은 기존 등록 데이터와 일치해야 합니다."],
    },
    "employees": {
        "filename": "직원_일괄등록_양식.xlsx",
        "permission": "menu.user.manage",
        "headers": ["사원번호", "이름", "재직상태", "입사일", "퇴사일", "이메일", "비고", "부서코드", "1번부서", "2번부서", "직급코드", "직급"],
        "guides": ["필수: 사원번호, 이름, 재직상태, 입사일, 퇴사일", "선택: 이메일, 비고, 부서코드, 1번부서, 2번부서, 직급코드, 직급", "부서/직급 값은 별도 DB 테이블에 저장됩니다.", "재직상태: 재직, 퇴사, 휴직", "날짜는 YYYY-MM-DD 형식으로 입력합니다."],
    },
}


def excel_cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return value.date().isoformat() if isinstance(value, datetime) else value.isoformat()
    return str(value).strip()


def decode_excel_upload(content: bytes) -> list[dict[str, str]]:
    if not content:
        raise HTTPException(status_code=400, detail="Excel 파일이 비어 있습니다.")
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        sheet = workbook["일괄등록"] if "일괄등록" in workbook.sheetnames else workbook.active
        values = sheet.iter_rows(values_only=True)
        header_row = next(values, None)
        if not header_row:
            raise HTTPException(status_code=400, detail="Excel 헤더를 확인해 주세요.")
        headers = [excel_cell_text(value) for value in header_row]
        rows = []
        for cells in values:
            row = {header: excel_cell_text(cells[index] if index < len(cells) else None) for index, header in enumerate(headers) if header}
            if any(row.values()):
                rows.append(row)
        return rows
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(status_code=400, detail="올바른 .xlsx 파일인지 확인해 주세요.") from exc


def csv_company(row: dict[str, str]) -> dict[str, Any]:
    return {
        "name": row.get("업체명", ""),
        "businessNo": row.get("사업자번호", ""),
        "representative": row.get("대표자", ""),
        "address": row.get("주소", ""),
        "manager": row.get("담당자명", ""),
        "phone": row.get("담당자연락처", ""),
        "email": row.get("담당자이메일", ""),
        "memo": row.get("비고", ""),
    }


def csv_employee(row: dict[str, str], refs: dict[str, Any]) -> dict[str, Any]:
    status_map = {"재직": "active", "퇴사": "resigned", "휴직": "leave", "active": "active", "resigned": "resigned", "leave": "leave"}
    organization_id = resolve_csv_employee_organization_id(row, refs)
    return {
        "employeeNo": row.get("사원번호", ""),
        "name": row.get("이름", ""),
        "employmentStatus": status_map.get(row.get("재직상태", ""), row.get("재직상태", "") or "active"),
        "isActive": status_map.get(row.get("재직상태", ""), row.get("재직상태", "") or "active") == "active",
        "joinedAt": row.get("입사일", ""),
        "resignedAt": row.get("퇴사일", ""),
        "email": row.get("이메일", ""),
        "memo": row.get("비고", ""),
        "departmentCode": row.get("부서코드", ""),
        "departmentParentName": row.get("1번부서", ""),
        "departmentName": row.get("2번부서", ""),
        "positionCode": row.get("직급코드", ""),
        "positionName": row.get("직급", ""),
        "organizationId": organization_id,
    }


def resolve_csv_employee_organization_id(row: dict[str, str], refs: dict[str, Any]) -> int | None:
    names = department_organization_candidates(row.get("1번부서", ""), row.get("2번부서", ""))
    for name in names:
        if name and name in refs["organizations_by_name"]:
            return refs["organizations_by_name"][name]
    for name in names:
        key = normalize_organization_name(name)
        if key and key in refs["organizations_by_normalized_name"]:
            return refs["organizations_by_normalized_name"][key]
    return None


def find_csv_company(row: dict[str, str], refs: dict[str, Any]) -> dict[str, Any] | None:
    business_no = row.get("업체사업자번호", "") or row.get("사업자번호", "")
    name = row.get("업체명", "")
    return refs["companies_by_business_no"].get(business_no) if business_no else refs["companies_by_name"].get(name)


def csv_contract(row: dict[str, str], refs: dict[str, Any]) -> dict[str, Any]:
    company = find_csv_company(row, refs)
    organization_name = row.get("담당부서", "")
    manager_username = row.get("담당자아이디", "")
    organization_id = refs["organizations_by_name"].get(organization_name)
    manager = refs["users_by_username"].get(manager_username)
    manager_user_id = manager and manager["id"]
    billing_map = {"일시불": "once", "3분할": "split3", "2분할": "split2", "월": "monthly", "월납": "monthly"}
    status_map = {"진행": "active", "검토": "pending", "만료": "expired", "종료": "closed"}
    auto_alert = row.get("자동알림", "예").lower() not in {"아니오", "n", "no", "false", "0"}
    return {
        "companyId": company and company["id"],
        "name": row.get("계약명", ""),
        "organizationId": organization_id,
        "managerUserId": manager_user_id,
        "startDate": row.get("시작일", ""),
        "endDate": row.get("종료일", ""),
        "amount": row.get("금액", "0").replace(",", ""),
        "billingCycle": billing_map.get(row.get("지급방법", ""), row.get("지급방법", "") or "once"),
        "status": status_map.get(row.get("상태", ""), row.get("상태", "") or "active"),
        "renewalDate": row.get("갱신일", ""),
        "alertDays": row.get("알림일수", "60"),
        "autoAlert": auto_alert,
        "memo": row.get("비고", ""),
    }


def csv_worklog(row: dict[str, str], refs: dict[str, Any]) -> dict[str, Any]:
    company = find_csv_company(row, refs)
    contract = refs["contracts_by_company_name"].get((company["id"], row.get("계약명"))) if company else None
    category_map = {"방문": "visit", "방문 점검": "visit", "원격": "remote", "원격지원": "remote", "장애": "incident", "장애 처리": "incident", "부품교체": "part", "부품 교체": "part", "정기점검": "regular"}
    status_map = {"접수": "open", "처리중": "progress", "진행중": "progress", "완료": "done", "보류": "hold"}
    return {
        "companyId": company and company["id"],
        "contractId": contract and contract["id"],
        "date": row.get("작업일", ""),
        "category": category_map.get(row.get("작업구분", ""), row.get("작업구분", "") or "visit"),
        "status": status_map.get(row.get("처리상태", ""), row.get("처리상태", "") or "done"),
        "worker": row.get("처리자", ""),
        "title": row.get("제목", ""),
        "content": row.get("작업내용", ""),
    }


def organization_for_payload(payload: dict[str, Any], user: dict[str, Any]) -> int | None:
    return parse_int(payload.get("organizationId")) or None


def ensure_valid_contract_assignment(payload: dict[str, Any], user: dict[str, Any]) -> None:
    role = user.get("role")
    payload_org_id = parse_int(payload.get("organizationId"))
    manager_user_id = parse_int(payload.get("managerUserId"))
    if role in FULL_ACCESS_ROLES:
        if payload_org_id and not is_selectable_contract_organization(payload_org_id):
            raise HTTPException(status_code=403, detail=CONTRACT_ASSIGNMENT_ERROR)
        return

    allowed_org_ids = {parse_int(value) for value in get_allowed_contract_organization_ids(user)}
    if not payload_org_id or payload_org_id not in allowed_org_ids:
        raise HTTPException(status_code=403, detail=CONTRACT_ASSIGNMENT_ERROR)
    if role == "team_member" and manager_user_id != parse_int(user.get("id")):
        raise HTTPException(status_code=403, detail=CONTRACT_ASSIGNMENT_ERROR)


def is_selectable_contract_organization(organization_id: int) -> bool:
    with db() as conn:
        with conn.cursor() as cur:
            cur.execute(
                """
                SELECT 1
                FROM app.organizations
                WHERE id = %s AND is_active = TRUE AND parent_id IS NOT NULL
                """,
                (organization_id,),
            )
            return cur.fetchone() is not None

def is_root_organization(organization_id: int) -> bool:
    with db() as conn:
        with conn.cursor() as cur:
            cur.execute(
                "SELECT parent_id FROM app.organizations WHERE id = %s AND is_active = TRUE",
                (organization_id,),
            )
            row = cur.fetchone()
    return bool(row and row.get("parent_id") is None)


def ensure_company_access(company_id: int, request: Request) -> None:
    if not company_id:
        raise HTTPException(status_code=404, detail="company not found")
    with db() as conn:
        with conn.cursor() as cur:
            cur.execute("SELECT 1 FROM public.companies WHERE id = %s", (company_id,))
            if not cur.fetchone():
                raise HTTPException(status_code=404, detail="company not found")


def ensure_contract_access(contract_id: int, request: Request) -> None:
    user = current_user_full(request)
    if user.get("role") in FULL_ACCESS_ROLES:
        return
    organization_ids = [parse_int(value) for value in get_allowed_contract_organization_ids(user)]
    organization_ids = [value for value in organization_ids if value]
    if not contract_id or not organization_ids:
        raise HTTPException(status_code=403, detail="조직 권한이 없습니다.")
    with db() as conn:
        with conn.cursor() as cur:
            cur.execute(
                """
                SELECT 1
                FROM app.contract_organizations
                WHERE contract_id = %s AND organization_id = ANY(%s)
                """,
                (contract_id, organization_ids),
            )
            if not cur.fetchone():
                raise HTTPException(status_code=403, detail="조직 권한이 없습니다.")

def ensure_document_access(document_id: int, request: Request) -> None:
    user = current_user_full(request)
    if user.get("role") in FULL_ACCESS_ROLES:
        return
    organization_ids = [parse_int(value) for value in get_allowed_contract_organization_ids(user)]
    organization_ids = [value for value in organization_ids if value]
    if not document_id or not organization_ids:
        raise HTTPException(status_code=403, detail="조직 권한이 없습니다.")
    with db() as conn:
        with conn.cursor() as cur:
            cur.execute(
                """
                SELECT 1
                FROM public.documents d
                LEFT JOIN app.contract_organizations co ON co.contract_id = d.contract_id
                LEFT JOIN app.document_organizations doo ON doo.document_id = d.id
                WHERE d.id = %s
                  AND (co.organization_id = ANY(%s) OR doo.organization_id = ANY(%s))
                """,
                (document_id, organization_ids, organization_ids),
            )
            if not cur.fetchone():
                raise HTTPException(status_code=403, detail="조직 권한이 없습니다.")

def ensure_item_company_access(table: str, item_id: int, request: Request) -> None:
    row = fetch_one(table, item_id)
    ensure_company_access(parse_int(row.get("companyId")), request)


def ensure_worklog_access(item_id: int, request: Request) -> None:
    user = current_user_full(request)
    if user.get("role") in FULL_ACCESS_ROLES:
        return
    organization_ids = [parse_int(value) for value in get_allowed_contract_organization_ids(user)]
    organization_ids = [value for value in organization_ids if value]
    if not item_id or not organization_ids:
        raise HTTPException(status_code=403, detail="유지보수 이력 조직 권한이 없습니다.")
    with db() as conn:
        with conn.cursor() as cur:
            cur.execute(
                """
                SELECT 1
                FROM work_logs w
                JOIN app.contract_organizations co ON co.contract_id = w.contract_id
                WHERE w.id = %s AND co.organization_id = ANY(%s)
                """,
                (item_id, organization_ids),
            )
            if not cur.fetchone():
                raise HTTPException(status_code=403, detail="유지보수 이력 조직 권한이 없습니다.")


def ensure_worklog_contract_access(payload: dict[str, Any], request: Request) -> None:
    user = current_user_full(request)
    if user.get("role") in FULL_ACCESS_ROLES:
        return
    contract_id = parse_int(payload.get("contractId"))
    if not contract_id:
        raise HTTPException(status_code=403, detail="관련 계약을 선택해 주세요.")
    ensure_contract_access(contract_id, request)
    company_id = parse_int(payload.get("companyId"))
    with db() as conn:
        with conn.cursor() as cur:
            cur.execute(
                "SELECT 1 FROM contracts WHERE id = %s AND company_id = %s",
                (contract_id, company_id),
            )
            if not cur.fetchone():
                raise HTTPException(status_code=400, detail="관련 계약과 업체가 일치하지 않습니다.")


def save_contract_organization(item_id: int, organization_id: int | None) -> None:
    with db() as conn:
        with conn.cursor() as cur:
            if organization_id:
                cur.execute(
                    """
                    INSERT INTO app.contract_organizations (contract_id, organization_id, updated_at)
                    VALUES (%s, %s, now())
                    ON CONFLICT (contract_id)
                    DO UPDATE SET organization_id = EXCLUDED.organization_id, updated_at = now()
                    """,
                    (item_id, organization_id),
                )
            else:
                cur.execute("DELETE FROM app.contract_organizations WHERE contract_id = %s", (item_id,))


def save_document_organization(item_id: int, organization_id: int | None) -> None:
    with db() as conn:
        with conn.cursor() as cur:
            if organization_id:
                cur.execute(
                    """
                    INSERT INTO app.document_organizations (document_id, organization_id, updated_at)
                    VALUES (%s, %s, now())
                    ON CONFLICT (document_id)
                    DO UPDATE SET organization_id = EXCLUDED.organization_id, updated_at = now()
                    """,
                    (item_id, organization_id),
                )
            else:
                cur.execute("DELETE FROM app.document_organizations WHERE document_id = %s", (item_id,))


def require_admin(request: Request) -> dict[str, Any]:
    user = getattr(request.state, "user", None)
    if not user or user.get("role") != "admin":
        raise HTTPException(status_code=403, detail="admin only")
    return user


def current_user(request: Request) -> dict[str, Any] | None:
    return getattr(request.state, "user", None)


def write_audit_log(request: Request, action: str, target_table: str, target_id: Any = None, detail: str = "") -> None:
    user = current_user(request) or {}
    with db() as conn:
        with conn.cursor() as cur:
            cur.execute(
                """
                INSERT INTO app.audit_logs (user_id, username, action, target_table, target_id, detail)
                VALUES (%s, %s, %s, %s, %s, %s)
                """,
                (
                    user.get("id"),
                    user.get("username"),
                    action,
                    target_table,
                    None if target_id is None else str(target_id),
                    detail[:1000] if detail else "",
                ),
            )


def smtp_enabled() -> bool:
    return bool(CONFIG.get("SMTP_SERVER") and CONFIG.get("SENDER_EMAIL"))


def send_email(to_email: str, subject: str, body: str) -> None:
    if not smtp_enabled():
        raise RuntimeError("SMTP 설정이 없습니다.")
    message = EmailMessage()
    message["From"] = CONFIG["SENDER_EMAIL"]
    message["To"] = to_email
    message["Subject"] = subject
    message.set_content(body)
    with smtplib.SMTP(CONFIG["SMTP_SERVER"], int(CONFIG.get("SMTP_PORT") or 25), timeout=15) as smtp:
        password = CONFIG.get("SENDER_PASSWORD") or ""
        if password:
            try:
                smtp.login(CONFIG["SENDER_EMAIL"], password)
            except smtplib.SMTPException:
                # 사내 SMTP 25번 포트가 인증 없이 허용되는 경우가 있어 인증 실패 시 무인증 발송을 한 번 허용합니다.
                pass
        smtp.send_message(message)


def fetch_contract_notification_targets(days: int) -> list[dict[str, Any]]:
    with db() as conn:
        with conn.cursor() as cur:
            cur.execute(
                """
                SELECT
                    c.id AS contract_id,
                    c.name AS contract_name,
                    c.end_date,
                    c.amount,
                    co.organization_id,
                    o.name AS organization_name,
                    comp.name AS company_name,
                    u.id AS user_id,
                    u.username,
                    u.display_name,
                    u.email
                FROM public.contracts c
                JOIN public.companies comp ON comp.id = c.company_id
                LEFT JOIN app.contract_organizations co ON co.contract_id = c.id
                LEFT JOIN app.organizations o ON o.id = co.organization_id
                LEFT JOIN app.contract_details cd ON cd.contract_id = c.id
                JOIN app.app_users u ON u.is_active = TRUE
                    AND NULLIF(u.email, '') IS NOT NULL
                    AND (
                        (cd.manager_user_id IS NOT NULL AND u.id = cd.manager_user_id)
                        OR (cd.manager_user_id IS NULL AND (u.role = 'admin' OR u.organization_id = co.organization_id))
                    )
                WHERE c.auto_alert IS DISTINCT FROM FALSE
                  AND c.status <> 'closed'
                  AND c.end_date = CURRENT_DATE + (%s * INTERVAL '1 day')
                  AND NOT EXISTS (
                      SELECT 1
                      FROM app.email_notification_logs l
                      WHERE l.contract_id = c.id
                        AND l.user_id = u.id
                        AND l.alert_days = %s
                  )
                ORDER BY c.end_date, comp.name, c.name, u.id
                """,
                (days, days),
            )
            return [row_to_client(row) for row in cur.fetchall()]


def build_contract_alert_email(target: dict[str, Any], days: int) -> tuple[str, str]:
    subject = f"[계약 만료 알림] {target.get('companyName')} - {target.get('contractName')} D-{days}"
    body = "\n".join(
        [
            "계약 만료 예정 알림입니다.",
            "",
            f"업체: {target.get('companyName') or '-'}",
            f"계약명: {target.get('contractName') or '-'}",
            f"담당부서: {target.get('organizationName') or '-'}",
            f"계약 종료일: {target.get('endDate') or '-'}",
            f"남은 기간: {days}일",
            "",
            "재계약 또는 종료 여부를 확인해 주세요.",
        ]
    )
    return subject, body


def record_email_notification(target: dict[str, Any], days: int, status: str, detail: str = "") -> None:
    with db() as conn:
        with conn.cursor() as cur:
            cur.execute(
                """
                INSERT INTO app.email_notification_logs
                    (contract_id, user_id, recipient_email, alert_days, status, detail)
                VALUES (%s, %s, %s, %s, %s, %s)
                ON CONFLICT (contract_id, user_id, alert_days) DO NOTHING
                """,
                (
                    parse_int(target.get("contractId")),
                    parse_int(target.get("userId")),
                    target.get("email"),
                    days,
                    status,
                    detail[:1000] if detail else "",
                ),
            )


def send_contract_expiry_notifications(force: bool = False) -> dict[str, Any]:
    if not smtp_enabled():
        return {"sent": 0, "failed": 0, "skipped": True, "detail": "SMTP 설정 없음"}
    sent = 0
    failed = 0
    details: list[str] = []
    for days in MAIL_ALERT_DAYS:
        for target in fetch_contract_notification_targets(days):
            try:
                subject, body = build_contract_alert_email(target, days)
                send_email(str(target.get("email") or ""), subject, body)
                record_email_notification(target, days, "sent")
                sent += 1
            except Exception as exc:
                failed += 1
                details.append(f"{target.get('email')}: {exc}")
                if force:
                    record_email_notification(target, days, "failed", str(exc))
    return {"sent": sent, "failed": failed, "skipped": False, "detail": "; ".join(details[:5])}


def mail_scheduler_loop() -> None:
    last_run: date | None = None
    while True:
        try:
            now = datetime.now()
            hour = int(CONFIG.get("MAIL_NOTIFY_HOUR") or 9)
            minute = int(CONFIG.get("MAIL_NOTIFY_MINUTE") or 0)
            if now.hour == hour and now.minute == minute and last_run != now.date():
                send_contract_expiry_notifications()
                last_run = now.date()
        except Exception:
            pass
        threading.Event().wait(60)


def start_mail_scheduler() -> None:
    global MAIL_SCHEDULER_STARTED
    if MAIL_SCHEDULER_STARTED:
        return
    MAIL_SCHEDULER_STARTED = True
    thread = threading.Thread(target=mail_scheduler_loop, daemon=True, name="contract-mail-scheduler")
    thread.start()


@app.middleware("http")
async def auth_middleware(request: Request, call_next):
    allow_paths = {"/api/health", "/api/auth/login", "/api/auth/me", "/api/auth/logout", "/api/auth/register", "/api/auth/register/check"}
    if request.url.path.startswith("/api/") and request.url.path not in allow_paths:
        user = read_session(request.cookies.get(SESSION_COOKIE))
        if not user:
            return JSONResponse({"detail": "login required"}, status_code=401)
        if user.get("role") == "team_member" and request.method == "DELETE":
            return JSONResponse({"detail": "팀원 권한은 삭제할 수 없습니다."}, status_code=403)
        request.state.user = user
    return await call_next(request)


@app.post("/api/auth/login")
def login(payload: dict[str, Any], response: Response) -> dict[str, Any]:
    username = str(payload.get("username") or "").strip()
    password = str(payload.get("password") or "")
    with db() as conn:
        with conn.cursor() as cur:
            cur.execute(
                """
                SELECT u.*, o.name AS organization_name, e.employee_no,
                       p.name AS position_name, e.position
                FROM app.app_users u
                LEFT JOIN app.organizations o ON o.id = u.organization_id
                LEFT JOIN app.employee_master e ON e.id = u.employee_id
                LEFT JOIN app.employee_positions p ON p.code = e.position_code
                WHERE u.username = %s AND u.is_active = TRUE
                  AND (u.employee_id IS NULL OR (e.is_active = TRUE AND e.employment_status = 'active'))
                """,
                (username,),
            )
            user = cur.fetchone()
    if not user or not verify_password(password, user["password_hash"]):
        raise HTTPException(status_code=401, detail="아이디 또는 비밀번호가 올바르지 않습니다.")
    response.set_cookie(
        SESSION_COOKIE,
        sign_session(int(user["id"]), user["username"], user["role"]),
        httponly=True,
        samesite="lax",
    )
    return {"user": client_user(user)}


@app.post("/api/auth/register/check")
def register_check(payload: dict[str, Any]) -> dict[str, Any]:
    name = str(payload.get("name") or "").strip()
    employee_no = str(payload.get("employeeNo") or "").strip()
    if not name or not employee_no:
        raise HTTPException(status_code=400, detail="이름과 사원번호를 입력해 주세요.")
    with db() as conn:
        with conn.cursor() as cur:
            cur.execute(
                """
                SELECT e.*, o.name AS organization_name, u.id AS user_id,
                       p.name AS position_name, e.position
                FROM app.employee_master e
                LEFT JOIN app.organizations o ON o.id = e.organization_id
                LEFT JOIN app.app_users u ON u.employee_id = e.id
                LEFT JOIN app.employee_positions p ON p.code = e.position_code
                WHERE e.employee_no = %s AND e.name = %s
                """,
                (employee_no, name),
            )
            employee = cur.fetchone()
    if not employee:
        raise HTTPException(status_code=404, detail="직원 기준정보와 일치하지 않습니다.")
    if not employee.get("is_active") or employee.get("employment_status") != "active":
        raise HTTPException(status_code=403, detail="재직 중인 직원만 가입할 수 있습니다.")
    if employee.get("user_id"):
        raise HTTPException(status_code=409, detail="이미 가입된 직원입니다.")
    return {
        "ok": True,
        "employee": {
            "employeeNo": employee.get("employee_no") or "",
            "name": employee.get("name") or "",
            "organizationName": employee.get("organization_name") or "",
            "email": employee.get("email") or "",
        },
    }


@app.post("/api/auth/register")
def register(payload: dict[str, Any], response: Response) -> dict[str, Any]:
    name = str(payload.get("name") or "").strip()
    employee_no = str(payload.get("employeeNo") or "").strip()
    username = employee_no
    email = str(payload.get("email") or "").strip()
    password = str(payload.get("password") or "")
    if not name or not employee_no or not password:
        raise HTTPException(status_code=400, detail="이름, 사원번호, 비밀번호는 필수입니다.")
    if not re.match(r"^[A-Za-z0-9_.-]{3,30}$", username):
        raise HTTPException(status_code=400, detail="사원번호는 영문, 숫자, ., _, - 조합 3~30자로 입력해 주세요.")
    if email and not re.match(r"^[^@\s]+@[^@\s]+\.[^@\s]+$", email):
        raise HTTPException(status_code=400, detail="이메일 형식이 올바르지 않습니다.")
    validate_signup_password(password)
    with db() as conn:
        with conn.cursor() as cur:
            cur.execute(
                """
                SELECT e.*, o.name AS organization_name, u.id AS user_id,
                       p.name AS position_name, e.position
                FROM app.employee_master e
                LEFT JOIN app.organizations o ON o.id = e.organization_id
                LEFT JOIN app.app_users u ON u.employee_id = e.id
                LEFT JOIN app.employee_positions p ON p.code = e.position_code
                WHERE e.employee_no = %s AND e.name = %s
                FOR UPDATE OF e
                """,
                (employee_no, name),
            )
            employee = cur.fetchone()
            if not employee:
                raise HTTPException(status_code=404, detail="직원 기준정보와 일치하지 않습니다.")
            if not employee.get("is_active") or employee.get("employment_status") != "active":
                raise HTTPException(status_code=403, detail="재직 중인 직원만 가입할 수 있습니다.")
            if employee.get("user_id"):
                raise HTTPException(status_code=409, detail="이미 가입된 직원입니다.")
            cur.execute("SELECT 1 FROM app.app_users WHERE username = %s", (username,))
            if cur.fetchone():
                raise HTTPException(status_code=409, detail="이미 사용 중인 아이디입니다.")
            organization_id = resolve_employee_organization_id(cur, employee)
            if organization_id and organization_id != employee.get("organization_id"):
                cur.execute(
                    "UPDATE app.employee_master SET organization_id = %s, updated_at = now() WHERE id = %s",
                    (organization_id, employee["id"]),
                )
            if email and not employee.get("email"):
                cur.execute(
                    "UPDATE app.employee_master SET email = %s, updated_at = now() WHERE id = %s",
                    (email, employee["id"]),
                )
            cur.execute(
                """
                INSERT INTO app.app_users
                    (username, password_hash, display_name, email, role, organization_id, employee_id, is_active, updated_at)
                VALUES (%s, %s, %s, %s, 'team_member', %s, %s, TRUE, now())
                RETURNING *
                """,
                (
                    username,
                    password_hash(password),
                    employee["name"],
                    email or employee.get("email"),
                    organization_id,
                    employee["id"],
                ),
            )
            created_id = cur.fetchone()["id"]
            cur.execute(
                """
                SELECT u.*, o.name AS organization_name, e.employee_no,
                       p.name AS position_name, e.position
                FROM app.app_users u
                LEFT JOIN app.organizations o ON o.id = u.organization_id
                LEFT JOIN app.employee_master e ON e.id = u.employee_id
                LEFT JOIN app.employee_positions p ON p.code = e.position_code
                WHERE u.id = %s
                """,
                (created_id,),
            )
            user = cur.fetchone()
    response.set_cookie(
        SESSION_COOKIE,
        sign_session(int(user["id"]), user["username"], user["role"]),
        httponly=True,
        samesite="lax",
    )
    return {"user": client_user(user)}


@app.post("/api/auth/logout")
def logout(response: Response) -> dict[str, bool]:
    response.delete_cookie(SESSION_COOKIE)
    return {"ok": True}


@app.get("/api/auth/me")
def auth_me(request: Request) -> dict[str, Any]:
    user = read_session(request.cookies.get(SESSION_COOKIE))
    if not user:
        return {"user": None}
    with db() as conn:
        with conn.cursor() as cur:
            cur.execute(
                """
                SELECT u.*, o.name AS organization_name, e.employee_no,
                       p.name AS position_name, e.position
                FROM app.app_users u
                LEFT JOIN app.organizations o ON o.id = u.organization_id
                LEFT JOIN app.employee_master e ON e.id = u.employee_id
                LEFT JOIN app.employee_positions p ON p.code = e.position_code
                WHERE u.id = %s AND u.is_active = TRUE
                  AND (u.employee_id IS NULL OR (e.is_active = TRUE AND e.employment_status = 'active'))
                """,
                (user["id"],),
            )
            row = cur.fetchone()
    return {"user": client_user(row) if row else None}


@app.get("/api/users")
def list_users(request: Request) -> dict[str, Any]:
    ensure_permission(request, "menu.user.view")
    employees = fetch_employees() if "menu.user.manage" in ROLE_PERMISSIONS.get(current_user_full(request).get("role") or "", set()) else []
    return {"users": fetch_users_for_admin(), "employees": employees}


@app.post("/api/users")
def create_user(payload: dict[str, Any], request: Request) -> dict[str, Any]:
    ensure_permission(request, "menu.user.manage")
    username = str(payload.get("username") or "").strip()
    password = str(payload.get("password") or "")
    if not username or not password:
        raise HTTPException(status_code=400, detail="아이디와 비밀번호는 필수입니다.")
    employee_id = parse_int(payload.get("employeeId")) or None
    with db() as conn:
        with conn.cursor() as cur:
            if employee_id:
                cur.execute("SELECT * FROM app.employee_master WHERE id = %s", (employee_id,))
                employee = cur.fetchone()
                if not employee:
                    raise HTTPException(status_code=400, detail="직원 기준정보를 찾을 수 없습니다.")
            cur.execute(
                """
                INSERT INTO app.app_users (username, password_hash, display_name, email, role, organization_id, employee_id, is_active, updated_at)
                VALUES (%s, %s, %s, %s, %s, %s, %s, %s, now())
                RETURNING *
                """,
                (
                    username,
                    password_hash(password),
                    payload.get("displayName"),
                    payload.get("email"),
                    payload.get("role") or "team_member",
                    parse_int(payload.get("organizationId")) or None,
                    employee_id,
                    parse_bool(payload.get("isActive", True)),
                ),
            )
            created_id = cur.fetchone()["id"]
            cur.execute(
                """
                SELECT u.*, o.name AS organization_name, e.employee_no,
                       p.name AS position_name, e.position
                FROM app.app_users u
                LEFT JOIN app.organizations o ON o.id = u.organization_id
                LEFT JOIN app.employee_master e ON e.id = u.employee_id
                LEFT JOIN app.employee_positions p ON p.code = e.position_code
                WHERE u.id = %s
                """,
                (created_id,),
            )
            created = client_user(cur.fetchone())
    write_audit_log(request, "create", "app.app_users", created["id"], created["username"])
    return created


@app.put("/api/users/{item_id}")
def update_user(item_id: int, payload: dict[str, Any], request: Request) -> dict[str, Any]:
    ensure_permission(request, "menu.user.manage")
    employee_id = parse_int(payload.get("employeeId")) or None
    values: list[Any] = [
        str(payload.get("username") or "").strip(),
        payload.get("displayName"),
        payload.get("email"),
        payload.get("role") or "team_member",
        parse_int(payload.get("organizationId")) or None,
        employee_id,
        parse_bool(payload.get("isActive", True)),
    ]
    password = str(payload.get("password") or "")
    password_sql = ""
    if password:
        password_sql = ", password_hash = %s"
        values.append(password_hash(password))
    values.append(item_id)
    with db() as conn:
        with conn.cursor() as cur:
            if employee_id:
                cur.execute("SELECT 1 FROM app.employee_master WHERE id = %s", (employee_id,))
                if not cur.fetchone():
                    raise HTTPException(status_code=400, detail="직원 기준정보를 찾을 수 없습니다.")
            cur.execute(
                f"""
                UPDATE app.app_users
                SET username = %s, display_name = %s, email = %s, role = %s, organization_id = %s, employee_id = %s, is_active = %s,
                    updated_at = now(){password_sql}
                WHERE id = %s
                RETURNING *
                """,
                tuple(values),
            )
            row = cur.fetchone()
            if not row:
                raise HTTPException(status_code=404, detail="not found")
            cur.execute(
                """
                SELECT u.*, o.name AS organization_name, e.employee_no,
                       p.name AS position_name, e.position
                FROM app.app_users u
                LEFT JOIN app.organizations o ON o.id = u.organization_id
                LEFT JOIN app.employee_master e ON e.id = u.employee_id
                LEFT JOIN app.employee_positions p ON p.code = e.position_code
                WHERE u.id = %s
                """,
                (item_id,),
            )
            updated = client_user(cur.fetchone())
    write_audit_log(request, "update", "app.app_users", item_id, updated["username"])
    return updated


@app.delete("/api/users/{item_id}")
def delete_user(item_id: int, request: Request) -> dict[str, bool]:
    user = ensure_permission(request, "menu.user.manage")
    if int(user["id"]) == item_id:
        raise HTTPException(status_code=400, detail="현재 로그인한 계정은 삭제할 수 없습니다.")
    with db() as conn:
        with conn.cursor() as cur:
            cur.execute("DELETE FROM app.app_users WHERE id = %s", (item_id,))
            if cur.rowcount == 0:
                raise HTTPException(status_code=404, detail="not found")
    write_audit_log(request, "delete", "app.app_users", item_id)
    return {"ok": True}


def employee_values(payload: dict[str, Any]) -> dict[str, Any]:
    status = payload.get("employmentStatus") or "active"
    if status not in {"active", "resigned", "leave"}:
        status = "active"
    return {
        "employee_no": str(payload.get("employeeNo") or "").strip(),
        "name": str(payload.get("name") or "").strip(),
        "department_code": str(payload.get("departmentCode") or "").strip(),
        "department_parent_name": str(payload.get("departmentParentName") or "").strip(),
        "department_name": str(payload.get("departmentName") or "").strip(),
        "position_code": str(payload.get("positionCode") or "").strip(),
        "position_name": str(payload.get("positionName") or "").strip(),
        "organization_id": parse_int(payload.get("organizationId")) or None,
        "position": str(payload.get("positionName") or "").strip(),
        "email": str(payload.get("email") or "").strip(),
        "phone": "",
        "employment_status": status,
        "is_active": status == "active",
        "joined_at": payload.get("joinedAt") or None,
        "resigned_at": payload.get("resignedAt") or None,
        "memo": str(payload.get("memo") or "").strip(),
    }


def validate_employee_values(values: dict[str, Any]) -> None:
    if not values["employee_no"] or not values["name"]:
        raise HTTPException(status_code=400, detail="사원번호와 이름은 필수입니다.")
    if values["employment_status"] not in {"active", "resigned", "leave"}:
        raise HTTPException(status_code=400, detail="재직상태는 재직, 퇴사, 휴직 중 하나여야 합니다.")
    if not values.get("joined_at") or not values.get("resigned_at"):
        raise HTTPException(status_code=400, detail="입사일과 퇴사일은 필수입니다.")
    for key, label_text in (("joined_at", "입사일"), ("resigned_at", "퇴사일")):
        try:
            date.fromisoformat(str(values[key]))
        except ValueError:
            raise HTTPException(status_code=400, detail=f"{label_text}은 YYYY-MM-DD 형식이어야 합니다.")


def upsert_employee_reference_tables(cur, values: dict[str, Any]) -> None:
    department_code = values.get("department_code") or ""
    department_name = values.get("department_name") or ""
    if department_code and department_name:
        parent_name = values.get("department_parent_name") or None
        raw_name = f"{parent_name or ''} {department_name}".strip()
        cur.execute(
            """
            INSERT INTO app.employee_departments (code, parent_name, name, raw_name, is_active, updated_at)
            VALUES (%s, %s, %s, %s, TRUE, now())
            ON CONFLICT (code)
            DO UPDATE SET parent_name = EXCLUDED.parent_name,
                          name = EXCLUDED.name,
                          raw_name = EXCLUDED.raw_name,
                          is_active = TRUE,
                          updated_at = now()
            """,
            (department_code, parent_name, department_name, raw_name),
        )
    position_code = values.get("position_code") or ""
    position_name = values.get("position_name") or ""
    if position_code and position_name:
        cur.execute(
            """
            INSERT INTO app.employee_positions (code, name, is_active, updated_at)
            VALUES (%s, %s, TRUE, now())
            ON CONFLICT (code)
            DO UPDATE SET name = EXCLUDED.name,
                          is_active = TRUE,
                          updated_at = now()
            """,
            (position_code, position_name),
        )


def sync_employee_user_status(cur, employee_id: int, values: dict[str, Any]) -> None:
    if values["employment_status"] != "active" or not values["is_active"]:
        cur.execute(
            """
            UPDATE app.app_users
            SET is_active = FALSE, updated_at = now()
            WHERE employee_id = %s
            """,
            (employee_id,),
        )


@app.get("/api/employees")
def list_employees(request: Request) -> dict[str, Any]:
    ensure_permission(request, "menu.user.manage")
    return {"employees": fetch_employees()}


@app.post("/api/employees")
def create_employee(payload: dict[str, Any], request: Request) -> dict[str, Any]:
    ensure_permission(request, "menu.user.manage")
    values = employee_values(payload)
    validate_employee_values(values)
    with db() as conn:
        with conn.cursor() as cur:
            upsert_employee_reference_tables(cur, values)
            cur.execute(
                """
                INSERT INTO app.employee_master
                    (employee_no, name, department_code, position_code, organization_id, position, email, phone, employment_status,
                     is_active, joined_at, resigned_at, memo, updated_at)
                VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, now())
                RETURNING id
                """,
                tuple(values[key] for key in (
                    "employee_no", "name", "department_code", "position_code", "organization_id", "position", "email", "phone",
                    "employment_status", "is_active", "joined_at", "resigned_at", "memo"
                )),
            )
            employee_id = cur.fetchone()["id"]
            sync_employee_user_status(cur, employee_id, values)
            cur.execute(
                """
                SELECT e.*, o.name AS organization_name,
                       d.parent_name AS department_parent_name, d.name AS department_name,
                       p.name AS position_name
                FROM app.employee_master e
                LEFT JOIN app.organizations o ON o.id = e.organization_id
                LEFT JOIN app.employee_departments d ON d.code = e.department_code
                LEFT JOIN app.employee_positions p ON p.code = e.position_code
                WHERE e.id = %s
                """,
                (employee_id,),
            )
            created = client_employee(cur.fetchone())
    write_audit_log(request, "create", "app.employee_master", created["id"], created["employeeNo"])
    return created


@app.put("/api/employees/{item_id}")
def update_employee(item_id: int, payload: dict[str, Any], request: Request) -> dict[str, Any]:
    ensure_permission(request, "menu.user.manage")
    values = employee_values(payload)
    validate_employee_values(values)
    with db() as conn:
        with conn.cursor() as cur:
            upsert_employee_reference_tables(cur, values)
            cur.execute(
                """
                UPDATE app.employee_master
                SET employee_no = %s, name = %s, department_code = %s, position_code = %s,
                    organization_id = %s, position = %s, email = %s, phone = %s,
                    employment_status = %s, is_active = %s, joined_at = %s, resigned_at = %s, memo = %s,
                    updated_at = now()
                WHERE id = %s
                RETURNING id
                """,
                (*tuple(values[key] for key in (
                    "employee_no", "name", "department_code", "position_code", "organization_id", "position", "email", "phone",
                    "employment_status", "is_active", "joined_at", "resigned_at", "memo"
                )), item_id),
            )
            if not cur.fetchone():
                raise HTTPException(status_code=404, detail="not found")
            sync_employee_user_status(cur, item_id, values)
            cur.execute(
                """
                SELECT e.*, o.name AS organization_name,
                       d.parent_name AS department_parent_name, d.name AS department_name,
                       p.name AS position_name
                FROM app.employee_master e
                LEFT JOIN app.organizations o ON o.id = e.organization_id
                LEFT JOIN app.employee_departments d ON d.code = e.department_code
                LEFT JOIN app.employee_positions p ON p.code = e.position_code
                WHERE e.id = %s
                """,
                (item_id,),
            )
            updated = client_employee(cur.fetchone())
    write_audit_log(request, "update", "app.employee_master", item_id, updated["employeeNo"])
    return updated


@app.delete("/api/employees/{item_id}")
def delete_employee(item_id: int, request: Request) -> dict[str, bool]:
    ensure_permission(request, "menu.user.manage")
    with db() as conn:
        with conn.cursor() as cur:
            cur.execute(
                """
                UPDATE app.employee_master
                SET is_active = FALSE, employment_status = 'resigned',
                    resigned_at = COALESCE(resigned_at, CURRENT_DATE), updated_at = now()
                WHERE id = %s
                RETURNING employee_no
                """,
                (item_id,),
            )
            row = cur.fetchone()
            if not row:
                raise HTTPException(status_code=404, detail="not found")
            cur.execute(
                """
                UPDATE app.app_users
                SET is_active = FALSE, updated_at = now()
                WHERE employee_id = %s
                """,
                (item_id,),
            )
    write_audit_log(request, "delete", "app.employee_master", item_id, row.get("employee_no") or "")
    return {"ok": True}


def organization_values(payload: dict[str, Any]) -> dict[str, Any]:
    return {
        "name": str(payload.get("name") or "").strip(),
        "parent_id": parse_int(payload.get("parentId")) or None,
        "sort_order": parse_int(payload.get("sortOrder")) or 0,
        "is_active": parse_bool(payload.get("isActive", True)),
    }


def fetch_organization(item_id: int) -> dict[str, Any]:
    with db() as conn:
        with conn.cursor() as cur:
            cur.execute("SELECT * FROM app.organizations WHERE id = %s", (item_id,))
            row = cur.fetchone()
            if not row:
                raise HTTPException(status_code=404, detail="not found")
            return row_to_client(row)


def ensure_valid_organization_payload(values: dict[str, Any], item_id: int | None = None) -> None:
    if not values["name"]:
        raise HTTPException(status_code=400, detail="조직명은 필수입니다.")
    parent_id = values.get("parent_id")
    if item_id and parent_id == item_id:
        raise HTTPException(status_code=400, detail="자기 자신을 상위 부서로 지정할 수 없습니다.")
    with db() as conn:
        with conn.cursor() as cur:
            if parent_id:
                cur.execute("SELECT 1 FROM app.organizations WHERE id = %s AND is_active = TRUE", (parent_id,))
                if not cur.fetchone():
                    raise HTTPException(status_code=400, detail="상위 부서를 찾을 수 없습니다.")
                if item_id:
                    cur.execute(
                        """
                        WITH RECURSIVE descendants AS (
                            SELECT id FROM app.organizations WHERE parent_id = %s
                            UNION ALL
                            SELECT o.id FROM app.organizations o
                            JOIN descendants d ON o.parent_id = d.id
                        )
                        SELECT 1 FROM descendants WHERE id = %s
                        """,
                        (item_id, parent_id),
                    )
                    if cur.fetchone():
                        raise HTTPException(status_code=400, detail="하위 부서를 상위 부서로 지정할 수 없습니다.")
            duplicate_sql = """
                SELECT 1
                FROM app.organizations
                WHERE name = %s
                  AND COALESCE(parent_id, 0) = COALESCE(%s, 0)
                  AND is_active = TRUE
            """
            duplicate_params: list[Any] = [values["name"], parent_id]
            if item_id:
                duplicate_sql += " AND id <> %s"
                duplicate_params.append(item_id)
            cur.execute(duplicate_sql, tuple(duplicate_params))
            if cur.fetchone():
                raise HTTPException(status_code=409, detail="같은 상위 부서에 동일한 조직명이 이미 있습니다.")


@app.get("/api/organizations")
def list_organizations(request: Request) -> dict[str, Any]:
    ensure_permission(request, "menu.organization.view")
    return {"organizations": fetch_organizations()}


@app.post("/api/organizations")
def create_organization(payload: dict[str, Any], request: Request) -> dict[str, Any]:
    ensure_permission(request, "menu.organization.manage")
    values = organization_values(payload)
    ensure_valid_organization_payload(values)
    try:
        with db() as conn:
            with conn.cursor() as cur:
                cur.execute(
                    """
                    INSERT INTO app.organizations (name, parent_id, sort_order, is_active, updated_at)
                    VALUES (%s, %s, %s, %s, now())
                    RETURNING *
                    """,
                    (values["name"], values["parent_id"], values["sort_order"], values["is_active"]),
                )
                created = row_to_client(cur.fetchone())
    except psycopg.errors.UniqueViolation as exc:
        raise HTTPException(status_code=409, detail="같은 상위 부서에 동일한 조직명이 이미 있습니다.") from exc
    write_audit_log(request, "create", "app.organizations", created["id"], created.get("name") or "")
    return created


@app.put("/api/organizations/{item_id}")
def update_organization(item_id: int, payload: dict[str, Any], request: Request) -> dict[str, Any]:
    ensure_permission(request, "menu.organization.manage")
    values = organization_values(payload)
    ensure_valid_organization_payload(values, item_id)
    try:
        with db() as conn:
            with conn.cursor() as cur:
                cur.execute(
                    """
                    UPDATE app.organizations
                    SET name = %s, parent_id = %s, sort_order = %s, is_active = %s, updated_at = now()
                    WHERE id = %s
                    RETURNING *
                    """,
                    (values["name"], values["parent_id"], values["sort_order"], values["is_active"], item_id),
                )
                row = cur.fetchone()
                if not row:
                    raise HTTPException(status_code=404, detail="not found")
                updated = row_to_client(row)
    except psycopg.errors.UniqueViolation as exc:
        raise HTTPException(status_code=409, detail="같은 상위 부서에 동일한 조직명이 이미 있습니다.") from exc
    write_audit_log(request, "update", "app.organizations", item_id, updated.get("name") or "")
    return updated


@app.delete("/api/organizations/{item_id}")
def delete_organization(item_id: int, request: Request) -> dict[str, bool]:
    ensure_permission(request, "menu.organization.manage")
    with db() as conn:
        with conn.cursor() as cur:
            cur.execute("SELECT 1 FROM app.organizations WHERE parent_id = %s AND is_active = TRUE", (item_id,))
            if cur.fetchone():
                raise HTTPException(status_code=409, detail="하위 부서가 있는 조직은 삭제할 수 없습니다.")
            checks = [
                ("app.app_users", "organization_id"),
                ("app.employee_master", "organization_id"),
                ("app.contract_organizations", "organization_id"),
                ("app.document_organizations", "organization_id"),
            ]
            for table, column in checks:
                cur.execute(f"SELECT 1 FROM {table} WHERE {column} = %s LIMIT 1", (item_id,))
                if cur.fetchone():
                    raise HTTPException(status_code=409, detail="사용 중인 조직은 삭제할 수 없습니다.")
            cur.execute("UPDATE app.organizations SET is_active = FALSE, updated_at = now() WHERE id = %s", (item_id,))
            if cur.rowcount == 0:
                raise HTTPException(status_code=404, detail="not found")
    write_audit_log(request, "delete", "app.organizations", item_id)
    return {"ok": True}


@app.put("/api/organizations/{item_id}/contract-links")
def update_organization_contract_links(item_id: int, payload: dict[str, Any], request: Request) -> dict[str, Any]:
    ensure_permission(request, "menu.organization.manage")
    target_ids = payload.get("organizationIds")
    if not isinstance(target_ids, list):
        raise HTTPException(status_code=400, detail="연계 조직 목록 형식이 올바르지 않습니다.")
    replace_organization_contract_links(item_id, target_ids)
    write_audit_log(request, "update_contract_links", "app.organizations", item_id, ",".join(map(str, target_ids)))
    return {"organizationContractLinks": fetch_organization_contract_links()}


@app.get("/api/health")
def health() -> dict[str, Any]:
    upload_dir = Path(CONFIG["UPLOAD_DIR"])
    with db() as conn:
        with conn.cursor() as cur:
            cur.execute("SELECT 1 AS ok")
            db_ok = cur.fetchone()["ok"] == 1
    return {
        "ok": True,
        "db": db_ok,
        "uploadDir": str(upload_dir),
        "uploadDirExists": upload_dir.exists(),
        "time": now_iso(),
    }


@app.get("/api/bootstrap")
def bootstrap(request: Request) -> dict[str, Any]:
    user = current_user_full(request)
    permissions = sorted(ROLE_PERMISSIONS.get(user.get("role") or "", set()))
    return {
        "companies": fetch_companies(user),
        "contracts": fetch_contracts(user),
        "assets": fetch_all("maintenance_assets"),
        "worklogs": fetch_worklogs(user),
        "pendingItems": fetch_pending_items(user),
        "documents": fetch_documents(user),
        "organizations": fetch_organizations(),
        "organizationContractLinks": fetch_organization_contract_links() if "menu.organization.manage" in permissions else [],
        "allowedContractOrganizationIds": get_allowed_contract_organization_ids(user),
        "users": fetch_selectable_users(user),
        "employees": fetch_employees() if "menu.user.manage" in permissions else [],
        "permissions": permissions,
        "config": {"defaultAlertDays": 60, "uploadDir": CONFIG["UPLOAD_DIR"]},
    }


@app.post("/api/ocr/company")
async def ocr_company(file: UploadFile = File(...)) -> dict[str, Any]:
    if not Path(CONFIG["TESSERACT_CMD"]).exists():
        raise HTTPException(status_code=500, detail="Tesseract OCR engine not found")
    suffix = Path(file.filename or "").suffix.lower()
    if suffix not in {".png", ".jpg", ".jpeg", ".bmp", ".tif", ".tiff", ".pdf"}:
        raise HTTPException(status_code=400, detail="이미지 또는 PDF 파일만 업로드할 수 있습니다.")

    with tempfile.NamedTemporaryFile(delete=False, suffix=suffix) as tmp:
        shutil.copyfileobj(file.file, tmp)
        tmp_path = Path(tmp.name)
    try:
        text = ocr_pdf(tmp_path) if suffix == ".pdf" else ocr_image(tmp_path)
        return {"fields": parse_company_ocr(text), "text": normalize_text(text)}
    finally:
        try:
            tmp_path.unlink(missing_ok=True)
        except OSError:
            pass


@app.post("/api/companies")
def create_company(payload: dict[str, Any], request: Request) -> dict[str, Any]:
    ensure_permission(request, "menu.company.create")
    ensure_company_not_duplicate(payload)
    primary_contact = primary_company_contact(payload)
    values = {
        "name": payload.get("name"),
        "business_no": payload.get("businessNo"),
        "address": payload.get("address"),
        "manager": primary_contact.get("name") or payload.get("manager"),
        "phone": primary_contact.get("phone") or payload.get("phone"),
        "email": primary_contact.get("email") or payload.get("email"),
        "memo": payload.get("memo"),
    }
    company = insert_row("companies", values)
    save_company_profile(int(company["id"]), payload)
    save_company_contacts(int(company["id"]), payload)
    result = fetch_company(int(company["id"]))
    write_audit_log(request, "create", "public.companies", result["id"], result.get("name") or "")
    return result


@app.put("/api/companies/{item_id}")
def update_company(item_id: int, payload: dict[str, Any], request: Request) -> dict[str, Any]:
    ensure_company_update_permission(request)
    ensure_company_access(item_id, request)
    ensure_company_not_duplicate(payload, item_id)
    primary_contact = primary_company_contact(payload)
    values = {
        "name": payload.get("name"),
        "business_no": payload.get("businessNo"),
        "address": payload.get("address"),
        "manager": primary_contact.get("name") or payload.get("manager"),
        "phone": primary_contact.get("phone") or payload.get("phone"),
        "email": primary_contact.get("email") or payload.get("email"),
        "memo": payload.get("memo"),
    }
    update_row("companies", item_id, values)
    save_company_profile(item_id, payload)
    save_company_contacts(item_id, payload)
    result = fetch_company(item_id)
    write_audit_log(request, "update", "public.companies", item_id, result.get("name") or "")
    return result


def fetch_company(item_id: int) -> dict[str, Any]:
    with db() as conn:
        with conn.cursor() as cur:
            cur.execute(
                """
                SELECT c.*, p.representative
                FROM companies c
                LEFT JOIN app.company_profiles p ON p.company_id = c.id
                WHERE c.id = %s
                """,
                (item_id,),
            )
            row = cur.fetchone()
            if not row:
                raise HTTPException(status_code=404, detail="not found")
            company = row_to_client(row)
            attach_company_contacts(cur, [company])
            return company


def save_company_profile(item_id: int, payload: dict[str, Any]) -> None:
    with db() as conn:
        with conn.cursor() as cur:
            cur.execute(
                """
                INSERT INTO app.company_profiles (company_id, representative, updated_at)
                VALUES (%s, %s, now())
                ON CONFLICT (company_id)
                DO UPDATE SET representative = EXCLUDED.representative, updated_at = now()
                """,
                (item_id, payload.get("representative")),
            )


def ensure_company_not_duplicate(payload: dict[str, Any], current_id: int | None = None) -> None:
    name = str(payload.get("name") or "").strip()
    business_no = str(payload.get("businessNo") or "").strip()
    if not name and not business_no:
        return
    with db() as conn:
        with conn.cursor() as cur:
            params: list[Any] = []
            if business_no:
                condition = "business_no = %s"
                params.append(business_no)
            else:
                condition = "lower(name) = lower(%s)"
                params.append(name)
            if current_id:
                condition += " AND id <> %s"
                params.append(current_id)
            cur.execute(f"SELECT id, name FROM public.companies WHERE {condition} LIMIT 1", tuple(params))
            existing = cur.fetchone()
            if existing:
                raise HTTPException(status_code=409, detail=f"이미 등록된 업체입니다: {existing['name']}")


def clean_company_contacts(payload: dict[str, Any]) -> list[dict[str, Any]]:
    raw_contacts = payload.get("contacts")
    if not isinstance(raw_contacts, list):
        raw_contacts = []
    contacts: list[dict[str, Any]] = []
    for index, raw_contact in enumerate(raw_contacts):
        if not isinstance(raw_contact, dict):
            continue
        contact = {
            "name": str(raw_contact.get("name") or "").strip(),
            "position": str(raw_contact.get("position") or "").strip(),
            "phone": str(raw_contact.get("phone") or "").strip(),
            "email": str(raw_contact.get("email") or "").strip(),
            "memo": str(raw_contact.get("memo") or "").strip(),
            "isPrimary": parse_bool(raw_contact.get("isPrimary")) or index == 0,
        }
        if any(contact[key] for key in ("name", "position", "phone", "email", "memo")):
            if not contact["name"]:
                contact["name"] = "담당자"
            contacts.append(contact)
    if not contacts:
        fallback = {
            "name": str(payload.get("manager") or "").strip(),
            "position": "",
            "phone": str(payload.get("phone") or "").strip(),
            "email": str(payload.get("email") or "").strip(),
            "memo": "",
            "isPrimary": True,
        }
        if any(fallback[key] for key in ("name", "phone", "email")):
            if not fallback["name"]:
                fallback["name"] = "담당자"
            contacts.append(fallback)
    if contacts and not any(contact["isPrimary"] for contact in contacts):
        contacts[0]["isPrimary"] = True
    primary_seen = False
    for contact in contacts:
        if contact["isPrimary"] and not primary_seen:
            primary_seen = True
        else:
            contact["isPrimary"] = False
    return contacts


def primary_company_contact(payload: dict[str, Any]) -> dict[str, Any]:
    contacts = clean_company_contacts(payload)
    if not contacts:
        return {}
    return next((contact for contact in contacts if contact.get("isPrimary")), contacts[0])


def save_company_contacts(item_id: int, payload: dict[str, Any]) -> None:
    contacts = clean_company_contacts(payload)
    with db() as conn:
        with conn.cursor() as cur:
            cur.execute("DELETE FROM app.company_contacts WHERE company_id = %s", (item_id,))
            for contact in contacts:
                cur.execute(
                    """
                    INSERT INTO app.company_contacts
                        (company_id, name, position, phone, email, memo, is_primary, updated_at)
                    VALUES (%s, %s, %s, %s, %s, %s, %s, now())
                    """,
                    (
                        item_id,
                        contact["name"],
                        contact["position"],
                        contact["phone"],
                        contact["email"],
                        contact["memo"],
                        contact["isPrimary"],
                    ),
                )


@app.delete("/api/companies/{item_id}")
def delete_company(item_id: int, request: Request) -> dict[str, bool]:
    ensure_permission(request, "menu.company.delete")
    ensure_company_access(item_id, request)
    result = delete_row("companies", item_id)
    write_audit_log(request, "delete", "public.companies", item_id)
    return result


@app.post("/api/contracts")
def create_contract(payload: dict[str, Any], request: Request) -> dict[str, Any]:
    user = ensure_contract_create_permission(request)
    ensure_valid_contract_assignment(payload, user)
    ensure_company_access(parse_int(payload.get("companyId")), request)
    result = insert_row("contracts", contract_values(payload))
    save_contract_organization(int(result["id"]), organization_for_payload(payload, user))
    save_contract_details(int(result["id"]), payload, user)
    result = fetch_contract(int(result["id"]))
    write_audit_log(request, "create", "public.contracts", result["id"], result.get("name") or "")
    return result


@app.put("/api/contracts/{item_id}")
def update_contract(item_id: int, payload: dict[str, Any], request: Request) -> dict[str, Any]:
    user = ensure_contract_update_permission(request)
    ensure_valid_contract_assignment(payload, user)
    ensure_contract_access(item_id, request)
    ensure_company_access(parse_int(payload.get("companyId")), request)
    result = update_row("contracts", item_id, contract_values(payload))
    save_contract_organization(item_id, organization_for_payload(payload, user))
    save_contract_details(item_id, payload, user)
    result = fetch_contract(item_id)
    write_audit_log(request, "update", "public.contracts", item_id, result.get("name") or "")
    return result


@app.delete("/api/contracts/{item_id}")
def delete_contract(item_id: int, request: Request) -> dict[str, bool]:
    ensure_contract_delete_permission(item_id, request)
    result = delete_row("contracts", item_id)
    write_audit_log(request, "delete", "public.contracts", item_id)
    return result


def fetch_contract(item_id: int) -> dict[str, Any]:
    with db() as conn:
        with conn.cursor() as cur:
            cur.execute(
                """
                SELECT c.*, co.organization_id, o.name AS organization_name, cd.department, cd.manager,
                       cd.manager_user_id,
                       COALESCE(NULLIF(mu.display_name, ''), mu.username) AS manager_user_name,
                       mu.email AS manager_user_email
                FROM contracts c
                LEFT JOIN app.contract_organizations co ON co.contract_id = c.id
                LEFT JOIN app.organizations o ON o.id = co.organization_id
                LEFT JOIN app.contract_details cd ON cd.contract_id = c.id
                LEFT JOIN app.app_users mu ON mu.id = cd.manager_user_id
                WHERE c.id = %s
                """,
                (item_id,),
            )
            row = cur.fetchone()
            if not row:
                raise HTTPException(status_code=404, detail="not found")
            return row_to_client(row)


def contract_values(payload: dict[str, Any]) -> dict[str, Any]:
    return {
        "company_id": parse_int(payload.get("companyId")),
        "name": payload.get("name"),
        "start_date": payload.get("startDate"),
        "end_date": payload.get("endDate"),
        "amount": parse_int(payload.get("amount")),
        "billing_cycle": payload.get("billingCycle") or "once",
        "status": payload.get("status") or "active",
        "renewal_date": payload.get("renewalDate") or None,
        "alert_days": parse_int(payload.get("alertDays"), 60) or 60,
        "auto_alert": parse_bool(payload.get("autoAlert", True)),
        "memo": payload.get("memo"),
    }


def ensure_valid_contract_manager_user(manager_user_id: int | None, user: dict[str, Any]) -> None:
    if not manager_user_id:
        return
    role = user.get("role")
    if role == "team_member" and manager_user_id != parse_int(user.get("id")):
        raise HTTPException(status_code=403, detail=CONTRACT_ASSIGNMENT_ERROR)
    with db() as conn:
        with conn.cursor() as cur:
            if role in FULL_ACCESS_ROLES:
                cur.execute("SELECT 1 FROM app.app_users WHERE id = %s AND is_active = TRUE", (manager_user_id,))
            elif role == "team_lead":
                cur.execute(
                    """
                    SELECT 1
                    FROM app.app_users
                    WHERE id = %s
                      AND is_active = TRUE
                      AND (
                          id = %s
                          OR (organization_id = %s AND role = 'team_member')
                      )
                    """,
                    (manager_user_id, parse_int(user.get("id")), parse_int(user.get("organizationId"))),
                )
            else:
                cur.execute("SELECT 1 WHERE FALSE")
            if not cur.fetchone():
                raise HTTPException(status_code=403, detail=CONTRACT_ASSIGNMENT_ERROR)


def save_contract_details(item_id: int, payload: dict[str, Any], user: dict[str, Any]) -> None:
    manager_user_id = parse_int(payload.get("managerUserId")) or None
    ensure_valid_contract_manager_user(manager_user_id, user)
    manager = str(payload.get("manager") or "").strip()
    with db() as conn:
        with conn.cursor() as cur:
            if manager_user_id:
                cur.execute(
                    "SELECT COALESCE(NULLIF(display_name, ''), username) AS manager_name FROM app.app_users WHERE id = %s",
                    (manager_user_id,),
                )
                manager_row = cur.fetchone()
                if manager_row:
                    manager = manager_row.get("manager_name") or manager
            elif "manager" not in payload:
                cur.execute("SELECT manager FROM app.contract_details WHERE contract_id = %s", (item_id,))
                existing = cur.fetchone()
                if existing:
                    manager = existing.get("manager") or ""
            cur.execute(
                """
                INSERT INTO app.contract_details (contract_id, department, manager, manager_user_id, updated_at)
                VALUES (%s, %s, %s, %s, now())
                ON CONFLICT (contract_id)
                DO UPDATE SET department = EXCLUDED.department,
                              manager = EXCLUDED.manager,
                              manager_user_id = EXCLUDED.manager_user_id,
                              updated_at = now()
                """,
                (item_id, payload.get("department"), manager, manager_user_id),
            )


@app.post("/api/pending-items")
def create_pending_item(payload: dict[str, Any], request: Request) -> dict[str, Any]:
    ensure_permission(request, "menu.pending.create")
    contract_id = parse_int(payload.get("contractId"))
    if contract_id:
        ensure_contract_access(contract_id, request)
    result = insert_row("app.pending_items", pending_item_values(payload))
    write_audit_log(request, "create", "app.pending_items", result["id"], result.get("title") or "")
    return result


@app.put("/api/pending-items/{item_id}")
def update_pending_item(item_id: int, payload: dict[str, Any], request: Request) -> dict[str, Any]:
    ensure_permission(request, "menu.pending.update")
    existing = fetch_one("app.pending_items", item_id)
    existing_contract_id = parse_int(existing.get("contractId"))
    if existing_contract_id:
        ensure_contract_access(existing_contract_id, request)
    contract_id = parse_int(payload.get("contractId"))
    if contract_id:
        ensure_contract_access(contract_id, request)
    result = update_row("app.pending_items", item_id, pending_item_values(payload))
    write_audit_log(request, "update", "app.pending_items", item_id, result.get("title") or "")
    return result


@app.delete("/api/pending-items/{item_id}")
def delete_pending_item(item_id: int, request: Request) -> dict[str, bool]:
    ensure_permission(request, "menu.pending.delete")
    existing = fetch_one("app.pending_items", item_id)
    contract_id = parse_int(existing.get("contractId"))
    if contract_id:
        ensure_contract_access(contract_id, request)
    result = delete_row("app.pending_items", item_id)
    write_audit_log(request, "delete", "app.pending_items", item_id)
    return result


def pending_item_values(payload: dict[str, Any]) -> dict[str, Any]:
    company_id = parse_int(payload.get("companyId"))
    contract_id = parse_int(payload.get("contractId"))
    status = payload.get("status") or "open"
    return {
        "company_id": company_id or None,
        "contract_id": contract_id or None,
        "title": payload.get("title"),
        "memo": payload.get("memo"),
        "due_date": payload.get("dueDate") or None,
        "status": status,
        "completed_at": datetime.now() if status == "done" else None,
    }


@app.post("/api/assets")
def create_asset(payload: dict[str, Any], request: Request) -> dict[str, Any]:
    ensure_permission(request, "menu.asset.create")
    ensure_company_access(parse_int(payload.get("companyId")), request)
    result = insert_row("maintenance_assets", asset_values(payload))
    write_audit_log(request, "create", "public.maintenance_assets", result["id"], result.get("name") or "")
    return result


@app.put("/api/assets/{item_id}")
def update_asset(item_id: int, payload: dict[str, Any], request: Request) -> dict[str, Any]:
    ensure_permission(request, "menu.asset.update")
    ensure_item_company_access("maintenance_assets", item_id, request)
    ensure_company_access(parse_int(payload.get("companyId")), request)
    result = update_row("maintenance_assets", item_id, asset_values(payload))
    write_audit_log(request, "update", "public.maintenance_assets", item_id, result.get("name") or "")
    return result


@app.delete("/api/assets/{item_id}")
def delete_asset(item_id: int, request: Request) -> dict[str, bool]:
    ensure_permission(request, "menu.asset.delete")
    ensure_item_company_access("maintenance_assets", item_id, request)
    result = delete_row("maintenance_assets", item_id)
    write_audit_log(request, "delete", "public.maintenance_assets", item_id)
    return result


def asset_values(payload: dict[str, Any]) -> dict[str, Any]:
    return {
        "company_id": parse_int(payload.get("companyId")),
        "type": payload.get("type") or "equipment",
        "name": payload.get("name"),
        "serial": payload.get("serial"),
        "location": payload.get("location"),
        "memo": payload.get("memo"),
    }


@app.post("/api/worklogs")
def create_worklog(payload: dict[str, Any], request: Request) -> dict[str, Any]:
    ensure_permission(request, "menu.worklog.create")
    ensure_company_access(parse_int(payload.get("companyId")), request)
    ensure_worklog_contract_access(payload, request)
    result = insert_row("work_logs", worklog_values(payload))
    write_audit_log(request, "create", "public.work_logs", result["id"], result.get("title") or "")
    return result


@app.put("/api/worklogs/{item_id}")
def update_worklog(item_id: int, payload: dict[str, Any], request: Request) -> dict[str, Any]:
    ensure_permission(request, "menu.worklog.update")
    ensure_worklog_access(item_id, request)
    ensure_company_access(parse_int(payload.get("companyId")), request)
    ensure_worklog_contract_access(payload, request)
    result = update_row("work_logs", item_id, worklog_values(payload))
    write_audit_log(request, "update", "public.work_logs", item_id, result.get("title") or "")
    return result


@app.delete("/api/worklogs/{item_id}")
def delete_worklog(item_id: int, request: Request) -> dict[str, bool]:
    ensure_permission(request, "menu.worklog.delete")
    ensure_worklog_access(item_id, request)
    result = delete_row("work_logs", item_id)
    write_audit_log(request, "delete", "public.work_logs", item_id)
    return result


def worklog_values(payload: dict[str, Any]) -> dict[str, Any]:
    contract_id = parse_int(payload.get("contractId"))
    return {
        "company_id": parse_int(payload.get("companyId")),
        "contract_id": contract_id or None,
        "work_date": payload.get("date"),
        "category": payload.get("category") or "visit",
        "status": payload.get("status") or "done",
        "worker": payload.get("worker"),
        "title": payload.get("title"),
        "content": payload.get("content"),
    }


def load_excel_import_refs() -> dict[str, Any]:
    with db() as conn:
        with conn.cursor() as cur:
            cur.execute("SELECT id, name, business_no FROM companies")
            companies = cur.fetchall()
            cur.execute("SELECT id, name FROM app.organizations WHERE is_active = TRUE AND parent_id IS NOT NULL")
            organizations = cur.fetchall()
            cur.execute("SELECT id, username, display_name FROM app.app_users WHERE is_active = TRUE")
            users = cur.fetchall()
            cur.execute("SELECT employee_no, name FROM app.employee_master")
            employees = cur.fetchall()
            cur.execute("SELECT id, company_id, name FROM contracts")
            contracts = cur.fetchall()
            cur.execute("SELECT contract_id, work_date, title FROM work_logs")
            worklogs = cur.fetchall()
    normalized_organizations: dict[str, list[int]] = {}
    for row in organizations:
        key = normalize_organization_name(row["name"])
        if key:
            normalized_organizations.setdefault(key, []).append(row["id"])
    unique_normalized_organizations = {
        key: ids[0]
        for key, ids in normalized_organizations.items()
        if len(set(ids)) == 1
    }
    return {
        "companies_by_name": {row["name"]: row for row in companies},
        "companies_by_business_no": {row["business_no"]: row for row in companies if row.get("business_no")},
        "company_names": {row["name"] for row in companies},
        "company_business_nos": {row["business_no"] for row in companies if row.get("business_no")},
        "organizations_by_name": {row["name"]: row["id"] for row in organizations},
        "organizations_by_normalized_name": unique_normalized_organizations,
        "users_by_username": {row["username"]: row for row in users},
        "employee_keys": {(row["employee_no"], row["name"]) for row in employees},
        "contracts_by_company_name": {(row["company_id"], row["name"]): row for row in contracts},
        "contract_keys": {(row["company_id"], row["name"]) for row in contracts},
        "worklog_keys": {(row["contract_id"], str(row["work_date"]), row["title"]) for row in worklogs},
    }


def validate_excel_rows(kind: str, rows: list[dict[str, str]]) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    payloads: list[dict[str, Any]] = []
    errors: list[dict[str, Any]] = []
    seen: set[tuple[Any, ...]] = set()
    seen_company_names: set[str] = set()
    seen_company_business_nos: set[str] = set()
    refs = load_excel_import_refs()
    for index, row in enumerate(rows, start=2):
        row_errors: list[str] = []
        payload: dict[str, Any]
        if kind == "companies":
            payload = csv_company(row)
            if not payload["name"]:
                row_errors.append("업체명은 필수입니다.")
            key = (payload["businessNo"] or "", payload["name"])
            if payload["name"] in refs["company_names"] or (payload["businessNo"] and payload["businessNo"] in refs["company_business_nos"]):
                row_errors.append("이미 등록된 업체명 또는 사업자번호입니다.")
            if payload["name"] in seen_company_names or (payload["businessNo"] and payload["businessNo"] in seen_company_business_nos):
                row_errors.append("Excel 파일 안에 업체명 또는 사업자번호가 중복되어 있습니다.")
            seen_company_names.add(payload["name"])
            if payload["businessNo"]:
                seen_company_business_nos.add(payload["businessNo"])
        elif kind == "employees":
            payload = csv_employee(row, refs)
            if not payload["employeeNo"] or not payload["name"]:
                row_errors.append("사원번호와 이름은 필수입니다.")
            if payload["employmentStatus"] not in {"active", "resigned", "leave"}:
                row_errors.append("재직상태는 재직, 퇴사, 휴직 중 하나여야 합니다.")
            for field, label_text in (("joinedAt", "입사일"), ("resignedAt", "퇴사일")):
                if not payload.get(field):
                    row_errors.append(f"{label_text}은 필수입니다.")
                    continue
                try:
                    date.fromisoformat(str(payload[field]))
                except ValueError:
                    row_errors.append(f"{label_text}은 YYYY-MM-DD 형식이어야 합니다.")
            key = (payload["employeeNo"], payload["name"])
            if key in refs["employee_keys"]:
                row_errors.append("이미 등록된 사원번호와 이름 조합입니다.")
        elif kind == "contracts":
            payload = csv_contract(row, refs)
            if not payload["companyId"]:
                row_errors.append("업체를 찾을 수 없습니다.")
            if not payload["name"]:
                row_errors.append("계약명은 필수입니다.")
            if not payload["organizationId"]:
                row_errors.append("담당부서를 찾을 수 없습니다.")
            if row.get("담당자아이디") and not payload["managerUserId"]:
                row_errors.append("담당자 아이디를 찾을 수 없습니다.")
            for field, label_text in (("startDate", "시작일"), ("endDate", "종료일")):
                try:
                    date.fromisoformat(str(payload[field]))
                except ValueError:
                    row_errors.append(f"{label_text}은 YYYY-MM-DD 형식이어야 합니다.")
            try:
                if date.fromisoformat(str(payload["endDate"])) < date.fromisoformat(str(payload["startDate"])):
                    row_errors.append("종료일은 시작일보다 빠를 수 없습니다.")
            except ValueError:
                pass
            if payload["billingCycle"] not in {"once", "split3", "split2", "monthly"}:
                row_errors.append("지급방법은 일시불, 3분할, 2분할, 월 중 하나여야 합니다.")
            if payload["status"] not in {"active", "pending", "expired", "closed"}:
                row_errors.append("상태는 진행, 검토, 만료, 종료 중 하나여야 합니다.")
            try:
                int(str(payload["amount"] or "0"))
                int(str(payload["alertDays"] or "60"))
            except ValueError:
                row_errors.append("금액과 알림일수는 숫자로 입력해야 합니다.")
            if payload["renewalDate"]:
                try:
                    date.fromisoformat(str(payload["renewalDate"]))
                except ValueError:
                    row_errors.append("갱신일은 YYYY-MM-DD 형식이어야 합니다.")
            key = (payload["companyId"], payload["name"])
            if key in refs["contract_keys"]:
                row_errors.append("동일 업체에 같은 계약명이 이미 등록되어 있습니다.")
        else:
            payload = csv_worklog(row, refs)
            if not payload["companyId"]:
                row_errors.append("업체를 찾을 수 없습니다.")
            if not payload["contractId"]:
                row_errors.append("해당 업체의 계약을 찾을 수 없습니다.")
            if not payload["title"]:
                row_errors.append("제목은 필수입니다.")
            try:
                date.fromisoformat(str(payload["date"]))
            except ValueError:
                row_errors.append("작업일은 YYYY-MM-DD 형식이어야 합니다.")
            if payload["category"] not in {"visit", "remote", "incident", "part", "regular"}:
                row_errors.append("작업구분 값을 확인해 주세요.")
            if payload["status"] not in {"open", "progress", "done", "hold"}:
                row_errors.append("처리상태 값을 확인해 주세요.")
            key = (payload["contractId"], payload["date"], payload["title"])
            if key in refs["worklog_keys"]:
                row_errors.append("같은 계약·작업일·제목의 이력이 이미 등록되어 있습니다.")
        if key in seen:
            row_errors.append("Excel 파일 안에 중복된 행이 있습니다.")
        seen.add(key)
        if row_errors:
            errors.append({"row": index, "messages": row_errors})
        payloads.append(payload)
    return payloads, errors


def insert_excel_rows(kind: str, payloads: list[dict[str, Any]]) -> int:
    with db() as conn:
        with conn.cursor() as cur:
            if kind == "companies":
                for payload in payloads:
                    cur.execute(
                        """
                        INSERT INTO companies (name, business_no, address, manager, phone, email, memo)
                        VALUES (%s, %s, %s, %s, %s, %s, %s)
                        RETURNING id
                        """,
                        (payload["name"], payload["businessNo"] or None, payload["address"], payload["manager"], payload["phone"], payload["email"], payload["memo"]),
                    )
                    company_id = cur.fetchone()["id"]
                    cur.execute(
                        "INSERT INTO app.company_profiles (company_id, representative, updated_at) VALUES (%s, %s, now())",
                        (company_id, payload["representative"]),
                    )
                    if payload["manager"]:
                        cur.execute(
                            """
                            INSERT INTO app.company_contacts
                                (company_id, name, phone, email, is_primary, updated_at)
                            VALUES (%s, %s, %s, %s, TRUE, now())
                            """,
                            (company_id, payload["manager"], payload["phone"], payload["email"]),
                        )
            elif kind == "employees":
                for payload in payloads:
                    values = employee_values(payload)
                    upsert_employee_reference_tables(cur, values)
                    cur.execute(
                        """
                        INSERT INTO app.employee_master
                            (employee_no, name, department_code, position_code, organization_id, position, email, phone, employment_status,
                             is_active, joined_at, resigned_at, memo, updated_at)
                        VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, now())
                        """,
                        tuple(values[key] for key in (
                            "employee_no", "name", "department_code", "position_code", "organization_id", "position", "email", "phone",
                            "employment_status", "is_active", "joined_at", "resigned_at", "memo"
                        )),
                    )
            elif kind == "contracts":
                cur.execute("SELECT id, username, display_name FROM app.app_users WHERE is_active = TRUE")
                manager_names = {
                    row["id"]: (row.get("display_name") or row["username"])
                    for row in cur.fetchall()
                }
                for payload in payloads:
                    values = contract_values(payload)
                    cur.execute(
                        """
                        INSERT INTO contracts
                            (company_id, name, start_date, end_date, amount, billing_cycle, status,
                             renewal_date, alert_days, auto_alert, memo)
                        VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
                        RETURNING id
                        """,
                        tuple(values[key] for key in (
                            "company_id", "name", "start_date", "end_date", "amount", "billing_cycle",
                            "status", "renewal_date", "alert_days", "auto_alert", "memo"
                        )),
                    )
                    contract_id = cur.fetchone()["id"]
                    cur.execute(
                        "INSERT INTO app.contract_organizations (contract_id, organization_id, updated_at) VALUES (%s, %s, now())",
                        (contract_id, payload["organizationId"]),
                    )
                    manager_user_id = parse_int(payload.get("managerUserId")) or None
                    cur.execute(
                        """
                        INSERT INTO app.contract_details
                            (contract_id, manager, manager_user_id, updated_at)
                        VALUES (%s, %s, %s, now())
                        """,
                        (contract_id, manager_names.get(manager_user_id, ""), manager_user_id),
                    )
            else:
                values = [worklog_values(payload) for payload in payloads]
                cur.executemany(
                    """
                    INSERT INTO work_logs
                        (company_id, contract_id, work_date, category, status, worker, title, content)
                    VALUES (%s, %s, %s, %s, %s, %s, %s, %s)
                    """,
                    [tuple(value[key] for key in (
                        "company_id", "contract_id", "work_date", "category", "status", "worker", "title", "content"
                    )) for value in values],
                )
    return len(payloads)


@app.get("/api/excel-template/{kind}")
def excel_template(kind: str, request: Request) -> Response:
    template = EXCEL_TEMPLATES.get(kind)
    if not template:
        raise HTTPException(status_code=404, detail="지원하지 않는 Excel 종류입니다.")
    ensure_permission(request, template["permission"])
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "일괄등록"
    sheet.append(template["headers"])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{sheet.cell(row=1, column=len(template['headers'])).coordinate}"
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(fill_type="solid", fgColor="0077C0")
    for index, header in enumerate(template["headers"], start=1):
        sheet.column_dimensions[sheet.cell(row=1, column=index).column_letter].width = max(14, min(28, len(header) * 2 + 4))
    guide = workbook.create_sheet("작성안내")
    guide.append(["Excel 일괄등록 작성안내"])
    guide["A1"].font = Font(bold=True, color="FFFFFF")
    guide["A1"].fill = PatternFill(fill_type="solid", fgColor="0077C0")
    for message in template["guides"]:
        guide.append([message])
    guide.column_dimensions["A"].width = 72
    output = io.BytesIO()
    workbook.save(output)
    filename = template["filename"]
    return Response(
        content=output.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(filename)}"},
    )


@app.post("/api/excel-import/{kind}")
async def excel_import(
    kind: str,
    request: Request,
    file: UploadFile = File(...),
    commit: bool = Form(False),
) -> dict[str, Any]:
    ensure_excel_import_admin(request)
    if kind not in {"companies", "contracts", "worklogs", "employees"}:
        raise HTTPException(status_code=404, detail="지원하지 않는 Excel 종류입니다.")
    if not str(file.filename or "").lower().endswith(".xlsx"):
        raise HTTPException(status_code=400, detail=".xlsx 파일만 업로드할 수 있습니다.")
    rows = decode_excel_upload(await file.read())
    if not rows:
        raise HTTPException(status_code=400, detail="등록할 데이터 행이 없습니다.")
    max_rows = 5000 if kind == "employees" else 1000
    if len(rows) > max_rows:
        raise HTTPException(status_code=400, detail=f"한 번에 최대 {max_rows:,}행까지 등록할 수 있습니다.")
    payloads, errors = validate_excel_rows(kind, rows)
    if errors or not commit:
        return {"total": len(rows), "valid": len(rows) - len(errors), "errors": errors, "committed": 0}
    created = insert_excel_rows(kind, payloads)
    audit_tables = {"companies": "public.companies", "contracts": "public.contracts", "worklogs": "public.work_logs", "employees": "app.employee_master"}
    write_audit_log(request, "excel_import", audit_tables[kind], None, f"{created} rows")
    return {"total": len(rows), "valid": len(rows), "errors": [], "committed": created}


@app.post("/api/documents")
async def create_document(
    request: Request,
    companyId: int = Form(...),
    contractId: int | None = Form(None),
    organizationId: int | None = Form(None),
    category: str = Form("contract"),
    title: str = Form(...),
    memo: str = Form(""),
    file: UploadFile | None = File(None),
) -> dict[str, Any]:
    ensure_permission(request, "menu.document.create")
    ensure_company_access(companyId, request)
    user = current_user_full(request)
    result = await save_document(None, companyId, contractId, organization_for_document(contractId, organizationId, user), category, title, memo, file)
    write_audit_log(request, "create", "public.documents", result["id"], result.get("title") or "")
    return result


@app.post("/api/documents/bulk")
async def create_documents_bulk(
    request: Request,
    companyId: int = Form(...),
    contractId: int | None = Form(None),
    organizationId: int | None = Form(None),
    category: str = Form("contract"),
    categories: list[str] | None = Form(None),
    memo: str = Form(""),
    files: list[UploadFile] = File(...),
) -> dict[str, Any]:
    ensure_permission(request, "menu.document.create")
    if not files:
        raise HTTPException(status_code=400, detail="업로드할 파일을 선택해 주세요.")
    if len(files) > 100:
        raise HTTPException(status_code=400, detail="한 번에 최대 100개 파일까지 등록할 수 있습니다.")
    ensure_company_access(companyId, request)
    user = current_user_full(request)
    resolved_organization_id = organization_for_document(contractId, organizationId, user)

    created: list[dict[str, Any]] = []
    errors: list[dict[str, Any]] = []
    for index, file in enumerate(files):
        filename = file.filename or "file"
        try:
            title = Path(filename).stem or filename
            file_category = categories[index] if categories and index < len(categories) and categories[index] else category
            result = await save_document(None, companyId, contractId, resolved_organization_id, file_category, title, memo, file)
            created.append(result)
            write_audit_log(request, "create", "public.documents", result["id"], result.get("title") or "")
        except Exception as exc:
            errors.append({"fileName": filename, "error": str(exc)})
    return {
        "total": len(files),
        "created": len(created),
        "failed": len(errors),
        "documents": created,
        "errors": errors,
    }


@app.put("/api/documents/{item_id}")
async def update_document(
    request: Request,
    item_id: int,
    companyId: int = Form(...),
    contractId: int | None = Form(None),
    organizationId: int | None = Form(None),
    category: str = Form("contract"),
    title: str = Form(...),
    memo: str = Form(""),
    file: UploadFile | None = File(None),
) -> dict[str, Any]:
    ensure_permission(request, "menu.document.update")
    ensure_document_access(item_id, request)
    ensure_company_access(companyId, request)
    user = current_user_full(request)
    result = await save_document(item_id, companyId, contractId, organization_for_document(contractId, organizationId, user), category, title, memo, file)
    write_audit_log(request, "update", "public.documents", item_id, result.get("title") or "")
    return result

async def save_document(
    item_id: int | None,
    company_id: int,
    contract_id: int | None,
    organization_id: int | None,
    category: str,
    title: str,
    memo: str,
    file: UploadFile | None,
) -> dict[str, Any]:
    values = {
        "company_id": company_id,
        "contract_id": contract_id or None,
        "category": category,
        "title": title,
        "memo": memo,
    }
    if file and file.filename:
        upload_root = Path(CONFIG["UPLOAD_DIR"])
        upload_root.mkdir(parents=True, exist_ok=True)
        month_dir = upload_root / datetime.now().strftime("%Y-%m")
        month_dir.mkdir(parents=True, exist_ok=True)
        stored_name = f"{datetime.now().strftime('%Y%m%d_%H%M%S')}_{uuid.uuid4().hex[:8]}_{clean_filename(file.filename)}"
        destination = month_dir / stored_name
        with destination.open("wb") as output:
            shutil.copyfileobj(file.file, output)
        values["file_name"] = file.filename
        values["file_path"] = str(destination)
    if item_id:
        result = update_row("documents", item_id, values)
        save_document_organization(item_id, organization_id)
        return fetch_document(item_id)
    result = insert_row("documents", values)
    save_document_organization(int(result["id"]), organization_id)
    return fetch_document(int(result["id"]))


def organization_for_document(contract_id: int | None, organization_id: int | None, user: dict[str, Any]) -> int | None:
    contract_id = parse_int(contract_id)
    if contract_id:
        if user.get("role") not in FULL_ACCESS_ROLES:
            ensure_contract_id_for_user(contract_id, user)
        with db() as conn:
            with conn.cursor() as cur:
                cur.execute("SELECT organization_id FROM app.contract_organizations WHERE contract_id = %s", (contract_id,))
                row = cur.fetchone()
                return int(row["organization_id"]) if row and row.get("organization_id") else None
    if user.get("role") in FULL_ACCESS_ROLES:
        return parse_int(organization_id) or None
    return parse_int(user.get("organizationId")) or None


def ensure_contract_id_for_user(contract_id: int, user: dict[str, Any]) -> None:
    organization_id = parse_int(user.get("organizationId"))
    with db() as conn:
        with conn.cursor() as cur:
            cur.execute(
                "SELECT 1 FROM app.contract_organizations WHERE contract_id = %s AND organization_id = %s",
                (contract_id, organization_id),
            )
            if not cur.fetchone():
                raise HTTPException(status_code=403, detail="조직 권한이 없습니다.")


def fetch_document(item_id: int) -> dict[str, Any]:
    with db() as conn:
        with conn.cursor() as cur:
            cur.execute(
                """
                SELECT d.*, doo.organization_id, o.name AS organization_name
                FROM documents d
                LEFT JOIN app.document_organizations doo ON doo.document_id = d.id
                LEFT JOIN app.organizations o ON o.id = doo.organization_id
                WHERE d.id = %s
                """,
                (item_id,),
            )
            row = cur.fetchone()
            if not row:
                raise HTTPException(status_code=404, detail="not found")
            return row_to_client(row)


@app.delete("/api/documents/{item_id}")
def delete_document(item_id: int, request: Request) -> dict[str, bool]:
    ensure_permission(request, "menu.document.delete")
    ensure_document_access(item_id, request)
    document = fetch_one("documents", item_id)
    result = delete_row("documents", item_id)
    file_path = document.get("filePath")
    if file_path:
        try:
            Path(file_path).unlink(missing_ok=True)
        except OSError:
            pass
    write_audit_log(request, "delete", "public.documents", item_id, document.get("title") or "")
    return result


@app.get("/api/documents/{item_id}/download")
def download_document(item_id: int, request: Request) -> FileResponse:
    ensure_document_access(item_id, request)
    document = fetch_one("documents", item_id)
    file_path = document.get("filePath")
    if not file_path or not Path(file_path).exists():
        raise HTTPException(status_code=404, detail="file not found")
    return FileResponse(file_path, filename=document.get("fileName") or Path(file_path).name)


@app.post("/api/documents/bulk-download")
def bulk_download_documents(payload: dict[str, Any], background_tasks: BackgroundTasks, request: Request) -> FileResponse:
    raw_ids = payload.get("ids") or []
    ids = [parse_int(item) for item in raw_ids if parse_int(item)]
    if not ids:
        raise HTTPException(status_code=400, detail="다운로드할 문서를 선택해 주세요.")
    for item_id in ids:
        ensure_document_access(item_id, request)

    placeholders = ", ".join(["%s"] * len(ids))
    with db() as conn:
        with conn.cursor() as cur:
            cur.execute(f"SELECT * FROM documents WHERE id IN ({placeholders}) ORDER BY id", tuple(ids))
            documents = [row_to_client(row) for row in cur.fetchall()]

    files: list[tuple[Path, str]] = []
    used_names: set[str] = set()
    for document in documents:
        file_path = document.get("filePath")
        if not file_path:
            continue
        path = Path(file_path)
        if not path.exists():
            continue
        original_name = clean_filename(document.get("fileName") or path.name)
        archive_name = unique_archive_name(original_name, used_names)
        files.append((path, archive_name))

    if not files:
        raise HTTPException(status_code=404, detail="다운로드 가능한 파일이 없습니다.")

    zip_path = Path(tempfile.gettempdir()) / f"documents_{datetime.now().strftime('%Y%m%d_%H%M%S')}_{uuid.uuid4().hex[:8]}.zip"
    with zipfile.ZipFile(zip_path, "w", zipfile.ZIP_DEFLATED) as archive:
        for path, archive_name in files:
            archive.write(path, archive_name)

    background_tasks.add_task(zip_path.unlink, missing_ok=True)
    return FileResponse(zip_path, filename=f"문서함_{datetime.now().strftime('%Y%m%d_%H%M%S')}.zip")


def unique_archive_name(name: str, used_names: set[str]) -> str:
    candidate = name
    stem = Path(name).stem
    suffix = Path(name).suffix
    index = 1
    while candidate.lower() in used_names:
        candidate = f"{stem} ({index}){suffix}"
        index += 1
    used_names.add(candidate.lower())
    return candidate


app.mount("/", StaticFiles(directory=BASE_DIR, html=True), name="static")
