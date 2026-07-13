from __future__ import annotations

import argparse
import re
import sys
from collections import Counter
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from tkinter import Tk, filedialog, messagebox

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

try:
    import psycopg
    from psycopg.rows import dict_row
except Exception:
    psycopg = None
    dict_row = None


LABELS = {
    "active": "재직",
    "resigned": "퇴사",
    "import": "일괄등록",
    "dept_ref": "부서분해_참고",
    "original": "원본",
    "emp_no": "사원번호",
    "name": "이름",
    "status": "재직상태",
    "join": "입사일",
    "resign": "퇴사일",
    "email": "이메일",
    "memo": "비고",
    "dept_code": "부서코드",
    "dept1": "1번부서",
    "dept2": "2번부서",
    "rank_code": "직급코드",
    "rank": "직급",
    "org_match": "조직매칭_확인",
    "system_org": "시스템조직명",
    "match_status": "매칭상태",
    "source_row": "원본행",
    "dedup": "중복조합 최신행 적용",
    "excluded": "중복조합 제외행",
}

IMPORT_HEADERS = [
    LABELS["emp_no"],
    LABELS["name"],
    LABELS["status"],
    LABELS["join"],
    LABELS["resign"],
    LABELS["email"],
    LABELS["memo"],
    LABELS["dept_code"],
    LABELS["dept1"],
    LABELS["dept2"],
    LABELS["rank_code"],
    LABELS["rank"],
]


@dataclass
class ConvertResult:
    output_path: Path
    total_rows: int
    active_rows: int
    resigned_rows: int
    duplicate_rows_removed: int
    organization_names_loaded: int
    auto_department_matches: int
    manual_department_matches: int


def clean(value) -> str:
    if value in (None, ""):
        return ""
    text = str(value).strip()
    if text.endswith(".0"):
        text = text[:-2]
    return text


def normalize_org_name(value: str) -> str:
    text = str(value or "").strip().lower()
    return re.sub(r"[\s\-_·.]+", "", text)


def department_org_candidates(parent_name: str, department_name: str) -> list[str]:
    candidates: list[str] = []
    department = str(department_name or "").strip()
    parent = str(parent_name or "").strip()
    if department:
        candidates.append(department)
        match = re.match(r"^(.+)\((하남|평동)\)$", department)
        if match:
            base_name = match.group(1).strip()
            site_name = match.group(2).strip()
            candidates.append(f"{site_name} {base_name}")
    if parent:
        candidates.append(parent)
    result: list[str] = []
    seen: set[str] = set()
    for candidate in candidates:
        key = normalize_org_name(candidate)
        if key and key not in seen:
            result.append(candidate)
            seen.add(key)
    return result


def load_config() -> dict[str, str]:
    config_path = Path(__file__).resolve().parent / "config.env"
    result: dict[str, str] = {}
    if not config_path.exists():
        return result
    for line in config_path.read_text(encoding="utf-8").splitlines():
        if "=" not in line or line.strip().startswith("#"):
            continue
        key, value = line.split("=", 1)
        result[key.strip()] = value.strip()
    return result


def load_system_organization_names() -> list[str]:
    if psycopg is None:
        return []
    config = load_config()
    if not config:
        return []
    conninfo = " ".join([
        f"host={config.get('DB_HOST', 'localhost')}",
        f"port={config.get('DB_PORT', '5432')}",
        f"dbname={config.get('DB_NAME', 'maintenance_db')}",
        f"user={config.get('DB_USER', 'maintenance_user')}",
        f"password={config.get('DB_PASSWORD', '')}",
    ])
    try:
        with psycopg.connect(conninfo, row_factory=dict_row) as conn:
            with conn.cursor() as cur:
                cur.execute(
                    """
                    SELECT name
                    FROM app.organizations
                    WHERE is_active = TRUE
                      AND parent_id IS NOT NULL
                    ORDER BY sort_order, id
                    """
                )
                return [str(row["name"]) for row in cur.fetchall()]
    except Exception:
        return []


def build_unique_org_index(organization_names: list[str]) -> dict[str, str]:
    grouped: dict[str, list[str]] = {}
    for name in organization_names:
        key = normalize_org_name(name)
        if key:
            grouped.setdefault(key, []).append(name)
    return {key: names[0] for key, names in grouped.items() if len(set(names)) == 1}


def normalize_department_for_system(parent_name: str, department_name: str, org_index: dict[str, str]) -> tuple[str, str, str, str]:
    if not org_index:
        return parent_name, department_name, "", "조직목록없음"

    candidates = department_org_candidates(parent_name, department_name)
    for index, candidate in enumerate(candidates):
        candidate_key = normalize_org_name(candidate)
        if candidate_key and candidate_key in org_index:
            matched = org_index[candidate_key]
            if index == 0:
                status = "자동보정" if matched != department_name else "일치"
                return parent_name, matched, matched, status
            if index == 1 and department_name and re.match(r"^(.+)\((하남|평동)\)$", department_name):
                return parent_name, matched, matched, "괄호조직자동보정"
            status = "상위부서일치" if matched == parent_name else "상위부서자동보정"
            return matched, department_name, matched, status

    return parent_name, department_name, "", "수동확인"


def ymd_to_date(value) -> date | None:
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip().replace("-", "").replace(".", "")
    if text.endswith(".0"):
        text = text[:-2]
    if len(text) != 8 or not text.isdigit():
        return None
    try:
        return date(int(text[:4]), int(text[4:6]), int(text[6:8]))
    except ValueError:
        return None


def split_department(raw) -> tuple[str, str]:
    text = str(raw or "").strip()
    if not text or text == "#N/A":
        return "", ""
    parts = text.split()
    if len(parts) >= 2:
        return parts[0], " ".join(parts[1:])
    return "", text


def build_department_map(ws, org_index: dict[str, str]) -> tuple[dict[str, tuple[str, str]], list[list[str]]]:
    result: dict[str, tuple[str, str]] = {}
    match_rows: list[list[str]] = []
    for row in ws.iter_rows(values_only=True):
        code = clean(row[0] if len(row) > 0 else "")
        if code:
            original_parent, original_name = split_department(row[1] if len(row) > 1 else "")
            parent_name, department_name, matched_org, status = normalize_department_for_system(original_parent, original_name, org_index)
            result[code] = (parent_name, department_name)
            match_rows.append([
                code,
                original_parent,
                original_name,
                parent_name,
                department_name,
                matched_org,
                status,
            ])
    return result, match_rows


def build_rank_map(ws) -> dict[str, str]:
    result: dict[str, str] = {}
    for row in ws.iter_rows(values_only=True):
        code = clean(row[0] if len(row) > 0 else "")
        name = str(row[1] or "").strip() if len(row) > 1 else ""
        if code:
            result[code] = "" if name == "#N/A" else name
    return result


def resolve_rank_code(raw_code: str, rank_map: dict[str, str]) -> str:
    if raw_code in rank_map:
        return raw_code
    if raw_code.isdigit() and len(raw_code) < 5:
        padded = raw_code.zfill(5)
        if padded in rank_map:
            return padded
    return raw_code


def make_output_path(source_path: Path, output_path: Path | None) -> Path:
    if output_path:
        return output_path
    return source_path.with_name("디케이_사원_명부_일괄등록용.xlsx")


def convert_roster(source_path: Path, output_path: Path | None = None) -> ConvertResult:
    if not source_path.exists():
        raise FileNotFoundError(f"원본 파일을 찾을 수 없습니다: {source_path}")

    workbook = load_workbook(source_path, data_only=True)
    if len(workbook.worksheets) < 3:
        raise ValueError("원본 엑셀에는 1번 시트, 부서코드 시트, 직급코드 시트가 필요합니다.")

    source_sheet = workbook.worksheets[0]
    organization_names = load_system_organization_names()
    org_index = build_unique_org_index(organization_names)
    dept_map, department_match_rows = build_department_map(workbook.worksheets[1], org_index)
    rank_map = build_rank_map(workbook.worksheets[2])

    records: dict[tuple[str, str], dict] = {}
    duplicate_rows_removed = 0

    for row_index in range(2, source_sheet.max_row + 1):
        employee_no = clean(source_sheet.cell(row_index, 1).value)
        name = clean(source_sheet.cell(row_index, 3).value)
        if not employee_no and not name:
            continue

        resigned_at = ymd_to_date(source_sheet.cell(row_index, 4).value)
        joined_at = ymd_to_date(source_sheet.cell(row_index, 6).value) or date(1900, 1, 1)
        status = LABELS["resigned"] if resigned_at else LABELS["active"]
        output_resigned_at = resigned_at or date(2099, 12, 31)

        department_code = clean(source_sheet.cell(row_index, 2).value)
        position_code = resolve_rank_code(clean(source_sheet.cell(row_index, 5).value), rank_map)
        department_parent, department_name = dept_map.get(department_code, ("", ""))
        position_name = rank_map.get(position_code, "")

        row = [
            employee_no,
            name,
            status,
            joined_at,
            output_resigned_at,
            "",
            f"{LABELS['source_row']}:{row_index}",
            department_code,
            department_parent,
            department_name,
            position_code,
            position_name,
        ]
        key = (employee_no, name)
        sort_key = (joined_at, output_resigned_at, row_index)
        record = {"sort_key": sort_key, "row": row}

        if key in records:
            duplicate_rows_removed += 1
            old = records[key]
            if sort_key >= old["sort_key"]:
                record["row"][6] += f" / {LABELS['dedup']}"
                records[key] = record
            else:
                old["row"][6] += f" / {LABELS['excluded']}:{row_index}"
        else:
            records[key] = record

    rows = [
        record["row"]
        for record in sorted(
            records.values(),
            key=lambda item: (
                item["row"][2] != LABELS["active"],
                item["row"][0],
                item["row"][1],
            ),
        )
    ]

    output = make_output_path(source_path, output_path)
    write_workbook(output, source_sheet, dept_map, rows, department_match_rows)

    status_counts = Counter(row[2] for row in rows)
    match_counts = Counter(row[6] for row in department_match_rows)
    return ConvertResult(
        output_path=output,
        total_rows=len(rows),
        active_rows=status_counts[LABELS["active"]],
        resigned_rows=status_counts[LABELS["resigned"]],
        duplicate_rows_removed=duplicate_rows_removed,
        organization_names_loaded=len(organization_names),
        auto_department_matches=sum(match_counts[key] for key in ("일치", "자동보정", "괄호조직자동보정", "상위부서일치", "상위부서자동보정")),
        manual_department_matches=match_counts["수동확인"],
    )


def write_workbook(output: Path, source_sheet, dept_map: dict[str, tuple[str, str]], rows: list[list], department_match_rows: list[list[str]]) -> None:
    workbook = Workbook()
    import_sheet = workbook.active
    import_sheet.title = LABELS["import"]
    import_sheet.append(IMPORT_HEADERS)
    for row in rows:
        import_sheet.append(row)

    dept_sheet = workbook.create_sheet(LABELS["dept_ref"])
    dept_sheet.append([LABELS["dept_code"], LABELS["dept1"], LABELS["dept2"]])
    for code, (parent_name, name) in sorted(dept_map.items()):
        dept_sheet.append([code, parent_name, name])

    match_sheet = workbook.create_sheet(LABELS["org_match"])
    match_sheet.append([
        LABELS["dept_code"],
        f"원본 {LABELS['dept1']}",
        f"원본 {LABELS['dept2']}",
        f"보정 {LABELS['dept1']}",
        f"보정 {LABELS['dept2']}",
        LABELS["system_org"],
        LABELS["match_status"],
    ])
    for row in department_match_rows:
        match_sheet.append(row)

    original_sheet = workbook.create_sheet(LABELS["original"])
    for row in source_sheet.iter_rows(values_only=True):
        original_sheet.append(list(row))

    style_workbook(workbook)
    for row in import_sheet.iter_rows(min_row=2, min_col=4, max_col=5):
        for cell in row:
            cell.number_format = "yyyy-mm-dd"

    workbook.save(output)


def style_workbook(workbook: Workbook) -> None:
    header_fill = PatternFill(fill_type="solid", fgColor="0077C0")
    header_font = Font(bold=True, color="FFFFFF")
    thin = Side(style="thin", color="C9DCEB")

    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        if sheet.max_row and sheet.max_column:
            sheet.auto_filter.ref = f"A1:{get_column_letter(sheet.max_column)}{sheet.max_row}"
        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for row in sheet.iter_rows():
            for cell in row:
                cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
                cell.alignment = Alignment(vertical="center")
        for col in range(1, sheet.max_column + 1):
            letter = get_column_letter(col)
            values = [cell.value for cell in sheet[letter][:120]]
            width = max(10, min(45, max((len(str(v)) for v in values if v is not None), default=8) + 2))
            sheet.column_dimensions[letter].width = width


def choose_file_with_dialog() -> Path | None:
    root = Tk()
    root.withdraw()
    root.attributes("-topmost", True)
    selected = filedialog.askopenfilename(
        title="사원명부 원본 엑셀 선택",
        filetypes=[("Excel files", "*.xlsx"), ("All files", "*.*")],
    )
    root.destroy()
    return Path(selected) if selected else None


def show_result(result: ConvertResult) -> None:
    message = (
        f"변환 완료\n\n"
        f"저장 위치: {result.output_path}\n"
        f"전체: {result.total_rows:,}건\n"
        f"재직: {result.active_rows:,}건\n"
        f"퇴사: {result.resigned_rows:,}건\n"
        f"중복 제거: {result.duplicate_rows_removed:,}건\n"
        f"시스템 조직명 로드: {result.organization_names_loaded:,}개\n"
        f"부서 자동보정/일치: {result.auto_department_matches:,}개\n"
        f"부서 수동확인 필요: {result.manual_department_matches:,}개"
    )
    try:
        root = Tk()
        root.withdraw()
        root.attributes("-topmost", True)
        messagebox.showinfo("사원명부 변환 완료", message)
        root.destroy()
    except Exception:
        print(message)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="디케이 사원명부를 직원 일괄등록 엑셀로 변환합니다.")
    parser.add_argument("source", nargs="?", help="원본 사원명부 엑셀 경로")
    parser.add_argument("-o", "--output", help="저장할 엑셀 경로")
    parser.add_argument("--no-dialog", action="store_true", help="파일 선택 창과 완료 메시지 창을 사용하지 않습니다.")
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    try:
        source = Path(args.source) if args.source else None
        if not source and not args.no_dialog:
            source = choose_file_with_dialog()
        if not source:
            print("원본 엑셀 파일을 선택하거나 경로를 지정해 주세요.")
            return 1
        output = Path(args.output) if args.output else None
        result = convert_roster(source, output)
        print(f"저장 위치: {result.output_path}")
        print(f"전체: {result.total_rows:,}건")
        print(f"재직: {result.active_rows:,}건")
        print(f"퇴사: {result.resigned_rows:,}건")
        print(f"중복 제거: {result.duplicate_rows_removed:,}건")
        print(f"시스템 조직명 로드: {result.organization_names_loaded:,}개")
        print(f"부서 자동보정/일치: {result.auto_department_matches:,}개")
        print(f"부서 수동확인 필요: {result.manual_department_matches:,}개")
        if not args.no_dialog:
            show_result(result)
        return 0
    except Exception as exc:
        if not args.no_dialog:
            try:
                root = Tk()
                root.withdraw()
                root.attributes("-topmost", True)
                messagebox.showerror("사원명부 변환 실패", str(exc))
                root.destroy()
            except Exception:
                pass
        print(f"오류: {exc}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
